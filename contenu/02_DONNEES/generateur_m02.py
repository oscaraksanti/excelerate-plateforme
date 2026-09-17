#!/usr/bin/env python3
# ══════════════════════════════════════════════════════════════════════
#  Données du module 2 — commissions BAOBAB et le budget saboté
#
#  Graine fixe. Formules en anglais avec préfixes _xlfn., comme Excel
#  les stocke réellement.
# ══════════════════════════════════════════════════════════════════════
import random, datetime as dt, pathlib, sys
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

random.seed(2026)
RACINE = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "sortie")
RACINE.mkdir(parents=True, exist_ok=True)

ENCRE = "FF101418"
Ftitre  = Font(name="Calibri", size=14, bold=True, color=ENCRE)
Fentete = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
Faide   = Font(name="Calibri", size=10, italic=True, color="FF6B7280")
Fgras   = Font(name="Calibri", size=11, bold=True)
Rentete = PatternFill("solid", fgColor="FF243044")
Rgris   = PatternFill("solid", fgColor="FFF2F4F6")
Rrouge  = PatternFill("solid", fgColor="FFFDECEA")
Bord    = Border(bottom=Side(style="thin", color="FFD6DBE1"))

def entete(ws, cols, ligne=1):
    for i, c in enumerate(cols, start=1):
        cel = ws.cell(ligne, i, c)
        cel.font, cel.fill = Fentete, Rentete
        cel.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[ligne].height = 22

def largeurs(ws, t):
    for i, w in enumerate(t, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def tableau(ws, nom, ref):
    t = Table(displayName=nom, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)

# ══════════════════════════════════════════════════════════════════════
#  1. Les 42 commerciaux
# ══════════════════════════════════════════════════════════════════════
AGENCES = [
    ("Kinshasa",    "CDF"), ("Lubumbashi", "CDF"), ("Abidjan", "XOF"),
    ("Dakar",       "XOF"), ("Douala",     "XAF"), ("Libreville", "XAF"),
]
PRENOMS = ["Nadège","Serge","Aïcha","Thomas","Aline","Junior","Grâce","Patrick",
           "Espérance","Dieudonné","Fatou","Moussa","Chantal","Blaise","Rachel",
           "Emmanuel","Joséphine","Célestin","Mariam","Olivier","Bernadette"]
NOMS = ["KALALA","KOUADIO","NDIAYE","MBARGA","ILUNGA","MUKENDI","LOKO","TSHIALA",
        "DIALLO","TRAORE","BAKAYOKO","NGOMA","MABIALA","SAMBA","OUEDRAOGO",
        "KOFFI","BINTOU","MASSAMBA","NZUZI","KIPRE","ESSONO"]

# Les taux du mois de référence — septembre 2026
TAUX = {"CDF": 3024, "XOF": 601, "XAF": 601}
PLAFOND_USD = 450

def realisation_cible(i):
    """Réparti volontairement autour des frontières de palier."""
    if i == 0:  return 1.000     # exactement à l'objectif
    if i == 1:  return 0.999     # à 0,1 point de l'objectif
    if i == 2:  return 1.200     # exactement à la frontière du dernier palier
    if i == 3:  return 0.000     # rien vendu
    if i == 4:  return 2.450     # très au-dessus : le plafond va mordre
    if i == 5:  return 0.800     # exactement au seuil d'entrée
    return round(random.uniform(0.55, 1.65), 3)

COMMERCIAUX = []
for i in range(42):
    agence, devise = AGENCES[i % 6]
    base = {"CDF": 9_000_000, "XOF": 1_800_000, "XAF": 1_800_000}[devise]
    objectif = int(round(base * random.uniform(0.7, 1.4) / 10_000) * 10_000)
    taux_r = realisation_cible(i)
    realise = int(round(objectif * taux_r / 1000) * 1000)

    if i == 6:      # objectif nul : la division par zéro, traitée dans l'énoncé
        objectif, realise = 0, int(base * 0.4)
    if i == 7:      # deux ans jour pour jour
        embauche = dt.date(2024, 9, 1)
    elif i == 8:    # un jour de moins que deux ans
        embauche = dt.date(2024, 9, 2)
    else:
        embauche = dt.date(random.randint(2016, 2026), random.randint(1, 12), random.randint(1, 28))
        if embauche > dt.date(2026, 8, 31):
            embauche = dt.date(2026, 8, 1)

    COMMERCIAUX.append({
        "matricule": f"BAO-{i+1:03d}",
        "nom": f"{random.choice(PRENOMS)} {random.choice(NOMS)}",
        "agence": agence, "devise": devise,
        "embauche": embauche, "objectif": objectif, "realise": realise,
    })

MOIS_REF = dt.date(2026, 9, 1)

# ── Le calcul de référence, en Python : il sert à vérifier Excel ──────
def annees_revolues(depuis: dt.date) -> int:
    """Années entières écoulées — exactement ce que rend DATEDIF(...;"y")."""
    n = MOIS_REF.year - depuis.year
    if (MOIS_REF.month, MOIS_REF.day) < (depuis.month, depuis.day):
        n -= 1
    return n

def taux_palier(objectif, realise):
    if objectif == 0: return 0.0
    r = realise / objectif
    if r < 0.80:  return 0.00
    if r < 1.00:  return 0.04
    if r < 1.20:  return 0.06
    return 0.08

def commission(c):
    t = taux_palier(c["objectif"], c["realise"])
    brute = t * c["realise"]
    atteint = c["objectif"] > 0 and c["realise"] >= c["objectif"]
    bonus = brute * 0.05 if (atteint and annees_revolues(c["embauche"]) >= 2) else 0.0
    totale = brute + bonus
    plafond_local = PLAFOND_USD * TAUX[c["devise"]]
    return min(totale, plafond_local), totale > plafond_local, bonus > 0

for c in COMMERCIAUX:
    c["commission"], c["plafonne"], c["bonus"] = commission(c)
    c["commission_usd"] = c["commission"] / TAUX[c["devise"]]

TOTAL_COMMISSION_USD = sum(c["commission_usd"] for c in COMMERCIAUX)
NB_OBJECTIF_ATTEINT = sum(1 for c in COMMERCIAUX if c["objectif"] > 0 and c["realise"] >= c["objectif"])
NB_BONUS = sum(1 for c in COMMERCIAUX if c["bonus"])
NB_PLAFONNES = sum(1 for c in COMMERCIAUX if c["plafonne"])

print(f"42 commerciaux · objectif atteint : {NB_OBJECTIF_ATTEINT} · bonus : {NB_BONUS} · plafonnés : {NB_PLAFONNES}")
print(f"commission totale : {TOTAL_COMMISSION_USD:,.2f} USD".replace(",", " "))

# ── Arrondis : sans eux, la somme de 42 divisions ne tombe jamais deux
#    fois pareil. La règle est écrite dans l'énoncé.
for c in COMMERCIAUX:
    c["commission"] = round(c["commission"])
    c["commission_usd"] = round(c["commission"] / TAUX[c["devise"]], 2)
    c["taux_realisation"] = 0.0 if c["objectif"] == 0 else c["realise"] / c["objectif"]
    c["anciennete"] = annees_revolues(c["embauche"])

def ecrire_commerciaux(ws, complet: bool):
    cols = ["Matricule", "Nom", "Agence", "Devise", "Date_Embauche", "Objectif", "Realise"]
    if complet:
        cols += ["Taux_Realisation", "Anciennete", "Taux_Commission",
                 "Commission_Locale", "Commission_USD"]
    entete(ws, cols)
    for i, c in enumerate(COMMERCIAUX, start=2):
        ws.cell(i, 1, c["matricule"]); ws.cell(i, 2, c["nom"])
        ws.cell(i, 3, c["agence"]);    ws.cell(i, 4, c["devise"])
        ws.cell(i, 5, c["embauche"]).number_format = "DD/MM/YYYY"
        ws.cell(i, 6, c["objectif"]);  ws.cell(i, 7, c["realise"])
        if complet:
            ws.cell(i, 8, round(c["taux_realisation"], 4)).number_format = "0,0 %"
            ws.cell(i, 9, c["anciennete"])
            ws.cell(i, 10, taux_palier(c["objectif"], c["realise"])).number_format = "0 %"
            ws.cell(i, 11, c["commission"])
            ws.cell(i, 12, c["commission_usd"]).number_format = "# ##0,00"
    largeurs(ws, [12, 26, 14, 9, 15, 14, 14, 16, 12, 16, 18, 16])
    ws.freeze_panes = "A2"

def ecrire_baremes(ws):
    ws["A1"] = "BAREMES — toutes les hypothèses, et nulle part ailleurs"; ws["A1"].font = Ftitre
    ws["A2"] = ("Une valeur d'hypothèse ne s'écrit jamais dans une formule. "
                "Elle vit ici, elle porte un nom, et tout le classeur pointe dessus.")
    ws["A2"].font = Faide
    entete(ws, ["Seuil bas", "Seuil haut", "Taux de commission"], ligne=4)
    for i, (bas, haut, taux) in enumerate(
            [(0, 0.80, 0.00), (0.80, 1.00, 0.04), (1.00, 1.20, 0.06), (1.20, 99, 0.08)], start=5):
        ws.cell(i, 1, bas).number_format = "0 %"
        ws.cell(i, 2, haut).number_format = "0 %"
        ws.cell(i, 3, taux).number_format = "0 %"
    tableau(ws, "t_Baremes", "A4:C8")

    ws["E4"] = "Bonus d'ancienneté";       ws["F4"] = 0.05
    ws["F4"].number_format = "0 %"
    ws["E5"] = "Ancienneté requise (ans)"; ws["F5"] = 2
    ws["E6"] = "Plafond de commission";    ws["F6"] = PLAFOND_USD
    ws["E7"] = "Mois de référence";        ws["F7"] = MOIS_REF
    ws["F7"].number_format = "MMM YYYY"
    for a in ("E4", "E5", "E6", "E7"): ws[a].font = Fgras

    for j, titre in ((5, "Devise"), (6, "Taux_USD")):
        cel = ws.cell(9, j, titre); cel.font, cel.fill = Fentete, Rentete
    ws.row_dimensions[9].height = 22
    for i, (d, t) in enumerate(TAUX.items(), start=10):
        ws.cell(i, 5, d); ws.cell(i, 6, t)
    tableau(ws, "t_Taux", "E9:F12")
    largeurs(ws, [14, 14, 20, 4, 26, 16])

# ══════════════════════════════════════════════════════════════════════
#  2. Le budget saboté — cinq anomalies, et cinq seulement
# ══════════════════════════════════════════════════════════════════════
BUDGET = [
    ("Kinshasa",   4_820_000, 3_140_000),
    ("Lubumbashi", 2_960_000, 2_105_000),
    ("Abidjan",    3_540_000, 2_390_000),
    ("Dakar",      2_780_000, 1_920_000),
    ("Douala",     3_110_000, 2_260_000),
    ("Libreville", 1_940_000, 1_505_000),
]
TAUX_HERITE = 3024
TAUX_PERIME = 2710
D9_EN_DUR  = 1_680_000   # au lieu de 3 540 000 − 2 390 000 = 1 150 000

def ecrire_herite(ws):
    """Le classeur que Nadège juge impossible. Il l'est.

    Cinq anomalies, chacune isolée dans sa propre cellule : mélangées,
    on ne saurait pas laquelle produit quel écart, et le TP deviendrait
    injuste.
    """
    ws["A1"] = "BUDGET 2025 — repris de l'ancien contrôleur"; ws["A1"].font = Ftitre
    ws["A2"] = "Nadège : « Ce chiffre est impossible. Tu as quinze minutes. »"
    ws["A2"].font = Faide
    ws["A4"] = "Taux CDF → USD"; ws["B4"] = TAUX_HERITE
    ws["A4"].font = Fgras

    entete(ws, ["Agence", "CA_CDF", "Charges_CDF", "Resultat_CDF", "Resultat_USD"], ligne=6)
    for i, (nom, ca, charges) in enumerate(BUDGET, start=7):
        ws.cell(i, 1, nom); ws.cell(i, 2, ca); ws.cell(i, 3, charges)
        # ── ANOMALIE 2 : une valeur en dur au milieu d'une colonne de formules
        ws.cell(i, 4, D9_EN_DUR if i == 9 else f"=B{i}-C{i}")
        # ── ANOMALIE 3 : un taux périmé écrit en dur au lieu de la cellule
        ws.cell(i, 5, f"=D{i}/{TAUX_PERIME}" if i == 10 else f"=D{i}/$B$4")

    ws.cell(14, 1, "TOTAL").font = Fgras
    # ── ANOMALIE 1 : la somme s'arrête deux lignes trop tôt
    ws.cell(14, 2, "=SUM(B7:B10)")
    ws.cell(14, 3, "=SUM(C7:C12)")
    ws.cell(14, 4, "=SUM(D7:D12)")
    ws.cell(14, 5, "=SUM(E7:E12)")
    for j in range(1, 6): ws.cell(14, j).fill = Rgris

    # ── ANOMALIE 4 : le contrôle général avale la ligne de total
    ws.cell(16, 1, "Contrôle général — détail vs total").font = Fgras
    ws.cell(16, 5, "=SUM(E7:E14)")
    ws.cell(16, 5).fill = Rrouge

    # ── ANOMALIE 5 : un lien vers le poste d'un collègue parti
    ws.cell(18, 1, "Report de l'exercice précédent").font = Fgras
    ws.cell(18, 2, "='C:\\Users\\ancien_collegue\\Budgets\\[Budget_2024.xlsx]Synthese'!$B$28")

    largeurs(ws, [36, 16, 16, 16, 16])

# ── Les valeurs de référence du fichier hérité, calculées ici ─────────
CA_CORRECT   = sum(ca for _, ca, _ in BUDGET)
CA_AFFICHE   = sum(ca for _, ca, _ in BUDGET[:4])
ECART_CA     = CA_CORRECT - CA_AFFICHE
RESULTAT_CORRECT_USD = round(sum(ca - ch for _, ca, ch in BUDGET) / TAUX_HERITE, 2)

print(f"budget : CA correct {CA_CORRECT:,} · affiché {CA_AFFICHE:,} · écart {ECART_CA:,}".replace(",", " "))
print(f"résultat consolidé correct : {RESULTAT_CORRECT_USD:,.2f} USD".replace(",", " "))

# ══════════════════════════════════════════════════════════════════════
#  3. Les feuilles de service
# ══════════════════════════════════════════════════════════════════════
def feuille_readme(wb, rempli=False):
    ws = wb.create_sheet("README")
    ws["A1"] = "README — la carte d'identité du classeur"; ws["A1"].font = Ftitre
    champs = [
        ("Objectif", "Calculer les commissions 2026 des 42 commerciaux et auditer le budget hérité."),
        ("Source des données", "Export RH du 1er septembre 2026 + Budget_2025 de l'ancien contrôleur."),
        ("Date de mise à jour", ""), ("Auteur", ""),
        ("Procédure d'actualisation", "Remplacer RAW, vérifier BAREMES, contrôler QUALITE."),
        ("Hypothèses", "Taux du mois de référence. Commission arrondie à l'unité en devise locale."),
        ("Définition — Taux de réalisation", "Réalisé ÷ Objectif. Vaut 0 si l'objectif est nul."),
        ("Définition — Commission", "Taux de palier × Réalisé, + bonus, plafonnée en USD."),
        ("Limites connues", ""), ("Contact", ""),
    ]
    for i, (k, v) in enumerate(champs, start=3):
        ws.cell(i, 1, k).font = Fgras
        ws.cell(i, 2, v)
        ws.cell(i, 1).border = Bord; ws.cell(i, 2).border = Bord
    largeurs(ws, [32, 84])

def feuille_qualite(wb, rempli=False):
    ws = wb.create_sheet("QUALITE")
    ws["A1"] = "QUALITE — les contrôles"; ws["A1"].font = Ftitre
    ws["A2"] = "Écris l'attendu avant de regarder le constaté. Sinon ce n'est plus un contrôle."
    ws["A2"].font = Faide
    entete(ws, ["Contrôle", "Attendu", "Constaté", "Statut"], ligne=4)
    for i, (c, a) in enumerate([
        ("Nombre de commerciaux", 42),
        ("Objectifs nuls (division par zéro à traiter)", 1),
        ("Commissions négatives", 0),
        ("Valeurs en dur restantes dans les formules", 0),
        ("Anomalies trouvées dans le fichier hérité", 5),
        ("Écart de contrôle : total − somme des devises", 0),
    ], start=5):
        ws.cell(i, 1, c); ws.cell(i, 2, a)
        if rempli: ws.cell(i, 3, a); ws.cell(i, 4, "PASS")
    largeurs(ws, [50, 12, 12, 10])

def feuille_annexe(wb):
    ws = wb.create_sheet("ANNEXE_IA")
    ws["A1"] = "ANNEXE_IA — ce que l'IA a fait, et ce qu'elle a raté"; ws["A1"].font = Ftitre
    ws["A2"] = ("Le prompt exact, la sortie, l'erreur trouvée, et le V du protocole qui l'a "
                "attrapée. « J'ai utilisé l'IA » ne vaut rien.")
    ws["A2"].font = Faide
    entete(ws, ["Étape", "Le prompt utilisé", "Ce que l'IA a rendu",
                "L'erreur détectée", "Le V qui l'a attrapée", "La correction"], ligne=4)
    ws.cell(5, 1, "Traduction de la règle de prime")
    ws.cell(6, 1, "Audit du fichier hérité")
    for i in range(5, 10):
        for j in range(1, 7): ws.cell(i, j).border = Bord
    largeurs(ws, [30, 46, 34, 34, 20, 40])

def feuille_audit(wb, rempli=False):
    ws = wb.create_sheet("AUDIT")
    ws["A1"] = "AUDIT — les cinq anomalies du fichier hérité"; ws["A1"].font = Ftitre
    ws["A2"] = ("Une par ligne. La cause, pas le symptôme : la cellule fausse n'est presque "
                "jamais celle qu'il faut corriger.")
    ws["A2"].font = Faide
    entete(ws, ["#", "Où", "Ce que j'ai constaté", "La cause",
                "L'impact chiffré", "La correction", "L'outil qui l'a trouvée"], ligne=4)
    for i in range(5, 10):
        ws.cell(i, 1, i - 4).font = Fgras
        for j in range(1, 8): ws.cell(i, j).border = Bord
    largeurs(ws, [5, 14, 42, 42, 20, 42, 30])

# ══════════════════════════════════════════════════════════════════════
#  4. LE TP 2
# ══════════════════════════════════════════════════════════════════════
REPONSES = [
    (4,  "Nombre de commerciaux ayant atteint leur objectif",
         '=COUNTIFS(t_Com[Taux_Realisation],">=1")'),
    (5,  "Nombre touchant le bonus d'ancienneté (objectif atteint ET ancienneté requise)",
         '=SUMPRODUCT((t_Com[Taux_Realisation]>=1)*(t_Com[Anciennete]>=AncienneteRequise))'),
    (6,  "Nombre de commerciaux dont la commission a été plafonnée",
         '=COUNTIFS(t_Com[Commission_USD],">="&PlafondUSD)'),
    (7,  "Commission totale versée, en USD",
         '=ROUND(SUM(t_Com[Commission_USD]),2)'),
    (8,  "CA total de l'agence de Kinshasa, en CDF",
         '=SUMIFS(t_Com[Realise],t_Com[Agence],"Kinshasa")'),
    (9,  "CA total des agences en zone XOF, converti en USD",
         '=SUMIFS(t_Com[Realise],t_Com[Devise],"XOF")/_xlfn.XLOOKUP("XOF",t_Taux[Devise],t_Taux[Taux_USD])'),
    (10, "Commission du matricule BAO-025, en devise locale",
         '=_xlfn.XLOOKUP("BAO-025",t_Com[Matricule],t_Com[Commission_Locale])'),
    (11, "Taux de réalisation moyen, toutes agences",
         '=AVERAGE(t_Com[Taux_Realisation])'),
    (12, "Nombre de commerciaux embauchés depuis au moins deux ans",
         '=COUNTIFS(t_Com[Date_Embauche],"<="&EDATE(MoisRef,-24))'),
    (13, "Montant oublié par la somme tronquée du fichier hérité, en CDF",
         '=SUM(HERITE!B11:B12)'),
    (14, "Résultat consolidé du fichier hérité, en USD, après correction",
         '=HERITE!E14'),
    (15, "Commission totale des agences en CDF, en USD",
         '=SUMIFS(t_Com[Commission_USD],t_Com[Devise],"CDF")'),
    (16, "Commission totale des agences en XOF, en USD",
         '=SUMIFS(t_Com[Commission_USD],t_Com[Devise],"XOF")'),
    (17, "Commission totale des agences en XAF, en USD",
         '=SUMIFS(t_Com[Commission_USD],t_Com[Devise],"XAF")'),
    (18, "Contrôle : total − somme des trois devises  →  doit valoir 0",
         '=ROUND(B7-(B15+B16+B17),2)'),
]

def herite_corrige(ws):
    """Le même budget, les cinq anomalies réparées."""
    ws["A1"] = "BUDGET 2025 — corrigé"; ws["A1"].font = Ftitre
    ws["A4"] = "Taux CDF → USD"; ws["B4"] = TAUX_HERITE; ws["A4"].font = Fgras
    entete(ws, ["Agence", "CA_CDF", "Charges_CDF", "Resultat_CDF", "Resultat_USD"], ligne=6)
    for i, (nom, ca, charges) in enumerate(BUDGET, start=7):
        ws.cell(i, 1, nom); ws.cell(i, 2, ca); ws.cell(i, 3, charges)
        ws.cell(i, 4, ca - charges)                       # valeurs : non notées
        ws.cell(i, 5, (ca - charges) / TAUX_HERITE)
    ws.cell(14, 1, "TOTAL").font = Fgras
    ws.cell(14, 2, "=SUM(B7:B12)"); ws.cell(14, 3, "=SUM(C7:C12)")
    ws.cell(14, 4, "=SUM(D7:D12)"); ws.cell(14, 5, "=SUM(E7:E12)")
    ws.cell(16, 1, "Contrôle général — détail vs total").font = Fgras
    ws.cell(16, 5, "=SUM(E7:E12)-E14")
    ws.cell(18, 1, "Report de l'exercice précédent — lien rompu").font = Fgras
    largeurs(ws, [36, 16, 16, 16, 16])

def tp02(corrige: bool):
    wb = openpyxl.Workbook()
    raw = wb.active; raw.title = "RAW"
    ecrire_commerciaux(raw, complet=False)

    com = wb.create_sheet("Commissions")
    if corrige:
        ecrire_commerciaux(com, complet=True)
        tableau(com, "t_Com", f"A1:L{len(COMMERCIAUX)+1}")
    else:
        entete(com, ["Matricule", "Nom", "Agence", "Devise", "Date_Embauche",
                     "Objectif", "Realise", "Taux_Realisation", "Anciennete",
                     "Taux_Commission", "Commission_Locale", "Commission_USD"])
        com["A3"] = ("Colle ici les 42 lignes de RAW, ajoute les cinq colonnes de calcul, "
                     "puis convertis en tableau structuré nommé t_Com (Ctrl+L).")
        com["A3"].font = Faide
        largeurs(com, [12, 26, 14, 9, 15, 14, 14, 16, 12, 16, 18, 16])

    ecrire_baremes(wb.create_sheet("BAREMES"))
    (herite_corrige if corrige else ecrire_herite)(wb.create_sheet("HERITE"))
    feuille_audit(wb, corrige)

    rep = wb.create_sheet("REPONSES")
    rep["A1"] = "REPONSES — une formule par cellule"; rep["A1"].font = Ftitre
    rep["A2"] = "Une valeur tapée à la main ne rapporte que la moitié des points."
    rep["A2"].font = Faide
    for ligne, libelle, formule in REPONSES:
        rep.cell(ligne, 1, libelle).border = Bord
        cel = rep.cell(ligne, 2); cel.border = Bord; cel.fill = Rgris
        if corrige: cel.value = formule
    largeurs(rep, [66, 22])

    NOMS = [("TauxBonus",         "BAREMES!$F$4"),
            ("AncienneteRequise", "BAREMES!$F$5"),
            ("PlafondUSD",        "BAREMES!$F$6"),
            ("MoisRef",           "BAREMES!$F$7")]
    if corrige:
        for nom, ref in NOMS:
            wb.defined_names.add(DefinedName(nom, attr_text=ref))
    else:
        bar = wb["BAREMES"]
        for i, (nom, _) in enumerate(NOMS, start=4):
            bar.cell(i, 7, f"← nomme cette cellule  {nom}").font = Faide
        largeurs(bar, [14, 14, 20, 4, 26, 16, 44])

    feuille_readme(wb, corrige); feuille_qualite(wb, corrige); feuille_annexe(wb)
    nom = "TP02_CORRIGE.xlsx" if corrige else "TP02_DEPART.xlsx"
    wb.save(RACINE / nom)

# ══════════════════════════════════════════════════════════════════════
#  5. Les classeurs des leçons
# ══════════════════════════════════════════════════════════════════════
def lecon_01():
    for suf, corrige in (("DEPART", False), ("CORRIGE", True)):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Commissions"
        ecrire_commerciaux(ws, complet=corrige)
        if corrige: tableau(ws, "t_Com", f"A1:L{len(COMMERCIAUX)+1}")
        ecrire_baremes(wb.create_sheet("BAREMES"))
        wb.save(RACINE / f"M02_L01_{suf}.xlsx")

def lecon_02():
    for suf, corrige in (("DEPART", False), ("CORRIGE", True)):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Commissions"
        ecrire_commerciaux(ws, complet=True)
        tableau(ws, "t_Com", f"A1:L{len(COMMERCIAUX)+1}")
        ecrire_baremes(wb.create_sheet("BAREMES"))
        r = wb.create_sheet("Reporting")
        r["A1"] = "Le reporting de Nadège — sans tableau croisé"; r["A1"].font = Ftitre
        r["A3"] = "Agence pilote"; r["B3"] = "Kinshasa"
        r["A5"] = "CA réalisé de l'agence pilote"
        r["A6"] = "Nombre de commerciaux de l'agence pilote"
        r["A7"] = "Commission moyenne de l'agence pilote (USD)"
        r["A8"] = "CA des commerciaux au-dessus de leur objectif"
        r["A9"] = "Total de contrôle — doit valoir 0"
        if corrige:
            r["B5"] = '=SUMIFS(t_Com[Realise],t_Com[Agence],$B$3)'
            r["B6"] = '=COUNTIFS(t_Com[Agence],$B$3)'
            r["B7"] = '=AVERAGEIFS(t_Com[Commission_USD],t_Com[Agence],$B$3)'
            r["B8"] = '=SUMPRODUCT((t_Com[Taux_Realisation]>=1)*t_Com[Realise])'
            r["B9"] = '=SUM(t_Com[Realise])-SUMPRODUCT(t_Com[Realise])'
        largeurs(r, [42, 20])
        wb.save(RACINE / f"M02_L02_{suf}.xlsx")

def lecon_03():
    for suf, corrige in (("DEPART", False), ("CORRIGE", True)):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Budget"
        (herite_corrige if corrige else ecrire_herite)(ws)
        wb.save(RACINE / f"M02_L03_{suf}.xlsx")
        if not corrige:
            wb2 = openpyxl.Workbook(); w = wb2.active; w.title = "Les 5 pieges"
            w["A1"] = "FICHIER FORMATEUR — ne jamais publier"; w["A1"].font = Ftitre
            entete(w, ["#", "Cellule", "L'anomalie", "L'outil qui la trouve", "Impact"], ligne=3)
            for i, l in enumerate([
                (1, "B14", "La somme s'arrête à B10 : Douala et Libreville sont oubliés",
                 "F5 > Différences entre lignes · comparer aux colonnes C et D",
                 f"{ECART_CA:,} CDF de CA manquants".replace(",", " ")),
                (2, "D9", "Une valeur en dur au milieu d'une colonne de formules",
                 "Triangle vert · F5 > Cellules > Formules",
                 "Résultat d'Abidjan surévalué de 530 000 CDF"),
                (3, "E10", "Le taux 2710 écrit en dur au lieu de $B$4",
                 "Ctrl+[ · rechercher 2710 dans les formules",
                 "Résultat de Dakar surévalué d'environ 11 %"),
                (4, "E16", "Le contrôle général somme aussi la ligne de total",
                 "Le total de contrôle n'est pas nul",
                 "Le consolidé apparaît deux fois"),
                (5, "B18", "Lien externe vers le poste d'un collègue parti",
                 "Données > Modifier les liens", "Valeur figée, jamais actualisée"),
            ], start=4):
                for j, v in enumerate(l, start=1): w.cell(i, j, v)
            largeurs(w, [5, 12, 56, 52, 44])
            wb2.save(RACINE / "M02_L03_FORMATEUR.xlsx")

if __name__ == "__main__":
    lecon_01(); lecon_02(); lecon_03()
    tp02(False); tp02(True)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Commerciaux"
    ecrire_commerciaux(ws, complet=False)
    wb.save(RACINE / "J02_Commissions_BAOBAB.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Budget"
    ecrire_herite(ws); wb.save(RACINE / "J02_Budget_SABOTE.xlsx")
    print("classeurs écrits dans", RACINE)
