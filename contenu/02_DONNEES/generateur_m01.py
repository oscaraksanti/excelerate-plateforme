#!/usr/bin/env python3
# ══════════════════════════════════════════════════════════════════════
#  Générateur des données du module 1 — GROUPE BAOBAB, agence de Kinshasa
#
#  Graine fixe : une régénération redonne EXACTEMENT les mêmes chiffres.
#  Sans ça, tous les corrigés deviennent faux du jour au lendemain.
#
#  Les formules sont écrites EN ANGLAIS avec les préfixes _xlfn. :
#  c'est ainsi qu'Excel les stocke réellement dans le fichier, quelle que
#  soit la langue de l'interface. Un Excel français les affichera en
#  français tout seul.
# ══════════════════════════════════════════════════════════════════════
import random, datetime as dt, pathlib, sys
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

GRAINE = 2026
random.seed(GRAINE)

RACINE = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "sortie")
RACINE.mkdir(parents=True, exist_ok=True)

# ── Palette de mise en forme ───────────────────────────────────────────
ENCRE   = "FF101418"
GRIS    = "FFF2F4F6"
VERT    = "FF2FBF71"
Ftitre  = Font(name="Calibri", size=14, bold=True, color=ENCRE)
Fentete = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
Faide   = Font(name="Calibri", size=10, italic=True, color="FF6B7280")
Fgras   = Font(name="Calibri", size=11, bold=True)
Rentete = PatternFill("solid", fgColor="FF243044")
Rgris   = PatternFill("solid", fgColor=GRIS)
Bord    = Border(bottom=Side(style="thin", color="FFD6DBE1"))

def entete(ws, colonnes, ligne=1):
    for i, c in enumerate(colonnes, start=1):
        cel = ws.cell(ligne, i, c)
        cel.font, cel.fill = Fentete, Rentete
        cel.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[ligne].height = 22

def largeurs(ws, tailles):
    for i, t in enumerate(tailles, start=1):
        ws.column_dimensions[get_column_letter(i)].width = t

def tableau(ws, nom, ref):
    t = Table(displayName=nom, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)

# ══════════════════════════════════════════════════════════════════════
#  1. L'univers
# ══════════════════════════════════════════════════════════════════════
CLIENTS = [
    "BOUTIQUE MAMA NGALULA", "SUPERMARCHE CITY", "DEPOT KINTAMBO",
    "ALIMENTATION BANDAL", "GROSSISTE MATONGE", "HOTEL FLEUVE CONGO",
    "RESTAURANT LE TROPICAL", "EPICERIE LEMBA", "SUPERETTE GOMBE",
    "DEPOT NGIRI-NGIRI", "BOUTIQUE SELEMBAO", "CANTINE UNIKIN",
    "MINI-MARCHE LIMETE", "ALIMENTATION KASA-VUBU", "GROSSISTE MASINA",
    "BOULANGERIE VICTOIRE", "PATISSERIE MA CAMPAGNE", "DEPOT KIMBANSEKE",
]

PRODUITS = [
    ("Huile végétale 5 L",        "Huiles",     34_000),
    ("Huile de palme 1 L",        "Huiles",      7_800),
    ("Riz parfumé 25 kg",         "Riz",       128_000),
    ("Riz local 50 kg",           "Riz",       165_000),
    ("Farine de froment 25 kg",   "Farine",     92_000),
    ("Farine de manioc 10 kg",    "Farine",     21_000),
    ("Eau minérale pack de 12",   "Boissons",   11_500),
    ("Jus d'ananas 1 L",          "Boissons",    4_200),
    ("Savon de ménage carton",    "Savons",     47_000),
    ("Savon de toilette lot 12",  "Savons",     18_600),
    ("Sardines boîte x50",        "Conserves",  62_000),
    ("Concentré de tomate x24",   "Conserves",  29_500),
]

COMMERCIAUX = ["Thomas MBARGA", "Aline KABEYA", "Junior ILUNGA",
               "Grâce MUKENDI", "Patrick LOKO"]

# ── La table des taux : 24 mois, octobre 2024 → septembre 2026 ─────────
TAUX = []
cdf, xof, xaf = 2_710, 596, 596
for k in range(24):
    mois = dt.date(2024, 10, 1)
    m = mois.month - 1 + k
    mois = dt.date(2024 + m // 12, m % 12 + 1, 1)
    cdf = round(cdf * random.uniform(1.000, 1.009))
    xof = round(xof + random.uniform(-4, 5))
    xaf = xof
    TAUX.append((mois, cdf, xof, xaf))

MOIS_ANALYSE = dt.date(2026, 9, 1)
TAUX_SEPT = next(c for m, c, _, _ in TAUX if m == MOIS_ANALYSE for c in [c])

# ══════════════════════════════════════════════════════════════════════
#  2. La caisse : 60 mouvements de septembre 2026
# ══════════════════════════════════════════════════════════════════════
LIGNES = []
for n in range(1, 61):
    jour = min(30, 1 + (n * 7 + random.randint(0, 3)) % 30)
    produit, famille, base = random.choice(PRODUITS)
    pu = int(round(base * random.uniform(0.94, 1.09) / 100) * 100)
    LIGNES.append({
        "date": dt.date(2026, 9, jour),
        "ticket": f"KIN-{n:04d}",
        "client": random.choice(CLIENTS),
        "produit": produit,
        "famille": famille,
        "quantite": random.randint(1, 38),
        "pu": pu,
        "commercial": random.choice(COMMERCIAUX),
    })
LIGNES.sort(key=lambda l: (l["date"], l["ticket"]))

# ── SUPERMARCHE CITY doit apparaître plusieurs fois, à des prix
#    différents : c'est ce qui rend la recherche « en marche arrière »
#    utile. Sans ça, premier prix et dernier prix sont identiques et la
#    pépite de la leçon 1.2 ne prouve rien.
positions_city = [3, 11, 24, 39, 52]
prix_city = [38_400, 41_200, 39_900, 42_700, 36_800]
for pos, prix in zip(positions_city, prix_city):
    LIGNES[pos]["client"] = "SUPERMARCHE CITY"
    LIGNES[pos]["produit"], LIGNES[pos]["famille"] = "Huile végétale 5 L", "Huiles"
    LIGNES[pos]["pu"] = prix

# ── Les six quantités stockées en texte : le piège du module ──────────
#    La somme d'une colonne qui en contient est fausse, et Excel ne dit
#    rien. C'est l'erreur la plus fréquente sur les exports comptables.
TEXTE_QTE = [5, 14, 22, 31, 43, 57]

DISTINCTS = sorted({l["client"] for l in LIGNES})
FAMILLES  = sorted({l["famille"] for l in LIGNES})

def ecrire_caisse(ws, propre: bool, avec_montant: bool):
    """La feuille de caisse. `propre` = quantités réparées."""
    entete(ws, ["Date", "Ticket", "Client", "Produit", "Famille",
                "Quantite", "PU_CDF", "Commercial"]
               + (["Montant"] if avec_montant else []))
    for i, l in enumerate(LIGNES, start=2):
        ws.cell(i, 1, l["date"]).number_format = "DD/MM/YYYY"
        ws.cell(i, 2, l["ticket"])
        ws.cell(i, 3, l["client"])
        ws.cell(i, 4, l["produit"])
        ws.cell(i, 5, l["famille"])
        q = l["quantite"]
        if not propre and (i - 2) in TEXTE_QTE:
            ws.cell(i, 6, str(q)).alignment = Alignment(horizontal="left")
        else:
            ws.cell(i, 6, q)
        ws.cell(i, 7, l["pu"]).number_format = '# ##0 "CDF"'
        ws.cell(i, 8, l["commercial"])
        if avec_montant:
            ws.cell(i, 9, q * l["pu"]).number_format = '# ##0 "CDF"'
    largeurs(ws, [12, 11, 26, 26, 12, 10, 14, 18, 16])
    ws.freeze_panes = "A2"

def ecrire_taux(ws):
    entete(ws, ["Mois", "CDF_USD", "XOF_USD", "XAF_USD"])
    for i, (mois, c, o, a) in enumerate(TAUX, start=2):
        ws.cell(i, 1, mois).number_format = "MMM YYYY"
        ws.cell(i, 2, c); ws.cell(i, 3, o); ws.cell(i, 4, a)
    tableau(ws, "t_Taux", f"A1:D{len(TAUX)+1}")
    largeurs(ws, [14, 12, 12, 12])

# ══════════════════════════════════════════════════════════════════════
#  3. La facture fournisseur — 14 articles
# ══════════════════════════════════════════════════════════════════════
FACTURE = [
    ("HV5-01", "Huile végétale 5 L",        8, 34_200),
    ("HP1-04", "Huile de palme 1 L",       36,  7_650),
    ("RZP-25", "Riz parfumé 25 kg",         6, 126_500),
    ("RZL-50", "Riz local 50 kg",           4, 162_000),
    ("FFR-25", "Farine de froment 25 kg",  10,  90_500),
    ("FMA-10", "Farine de manioc 10 kg",   18,  20_400),
    ("EAU-12", "Eau minérale pack de 12",  24,  11_200),
    ("JAN-01", "Jus d'ananas 1 L",         48,   4_050),
    ("SVM-CT", "Savon de ménage carton",    5,  46_300),
    ("SVT-12", "Savon de toilette lot 12", 14,  18_200),
    ("SRD-50", "Sardines boîte x50",        7,  61_400),
    ("CTM-24", "Concentré de tomate x24",  12,  28_900),
    ("HV5-02", "Huile végétale 2 L",       15,  14_700),
    ("RZP-10", "Riz parfumé 10 kg",         9,  53_800),
]
TOTAL_FACTURE = sum(q * p for _, _, q, p in FACTURE)
NUM_FACTURE = "FN-2026-0914"
DATE_FACTURE = dt.date(2026, 9, 14)

print(f"taux septembre 2026 : {TAUX_SEPT} CDF/USD")
print(f"total facture       : {TOTAL_FACTURE:,} CDF".replace(",", " "))
print(f"clients distincts   : {len(DISTINCTS)}")
print(f"familles            : {';'.join(FAMILLES)}")

# ══════════════════════════════════════════════════════════════════════
#  4. Les feuilles de service, communes aux classeurs de TP
# ══════════════════════════════════════════════════════════════════════
def feuille_readme(wb, rempli=False):
    ws = wb.create_sheet("README")
    ws["A1"] = "README — la carte d'identité du classeur"; ws["A1"].font = Ftitre
    ws["A2"] = "Sans cette feuille, personne ne peut reprendre ton fichier. Pas même toi dans six mois."
    ws["A2"].font = Faide
    champs = [
        ("Objectif du classeur", "Rendre exploitable la caisse de septembre 2026 de l'agence de Kinshasa."),
        ("Source des données", "Fichier de caisse tenu à la main par Thomas MBARGA + facture fournisseur FN-2026-0914."),
        ("Date de mise à jour", "à compléter"),
        ("Auteur", "à compléter"),
        ("Procédure d'actualisation", "Coller le nouvel export dans RAW, recopier dans Caisse, vérifier QUALITE."),
        ("Hypothèses", "Taux de change du 1er du mois. Montant = Quantité × PU_CDF. Aucune remise."),
        ("Définition — CA total", "Somme des montants en CDF, TVA non applicable sur ces flux de caisse."),
        ("Définition — Panier moyen", "CA total divisé par le nombre de tickets."),
        ("Limites connues", "à compléter"),
        ("Contact", "à compléter"),
    ]
    for i, (k, v) in enumerate(champs, start=4):
        ws.cell(i, 1, k).font = Fgras
        ws.cell(i, 2, v if rempli else ("" if "compléter" in v else v))
        ws.cell(i, 1).border = Bord; ws.cell(i, 2).border = Bord
    largeurs(ws, [30, 82])
    return ws

def feuille_qualite(wb, nb_lignes, ca_total, rempli=False):
    ws = wb.create_sheet("QUALITE")
    ws["A1"] = "QUALITE — les contrôles"; ws["A1"].font = Ftitre
    ws["A2"] = "Un contrôle qui n'est pas FAIL ne prouve rien s'il n'a jamais pu l'être. Écris l'attendu avant de regarder le constaté."
    ws["A2"].font = Faide
    entete(ws, ["Contrôle", "Attendu", "Constaté", "Statut"], ligne=4)
    controles = [
        ("Nombre de lignes reprises depuis RAW", nb_lignes),
        ("Quantités stockées en texte restantes", 0),
        ("Tickets en double", 0),
        ("Quantités négatives ou nulles", 0),
        ("Écart facture : total recalculé − total imprimé", 0),
        ("Écart de contrôle : lignes Caisse − lignes RAW", 0),
    ]
    for i, (c, a) in enumerate(controles, start=5):
        ws.cell(i, 1, c); ws.cell(i, 2, a)
        if rempli:
            ws.cell(i, 3, a); ws.cell(i, 4, "PASS")
    largeurs(ws, [46, 12, 12, 10])
    return ws

def feuille_annexe(wb):
    ws = wb.create_sheet("ANNEXE_IA")
    ws["A1"] = "ANNEXE_IA — ce que l'IA a fait, et ce qu'elle a raté"; ws["A1"].font = Ftitre
    ws["A2"] = ("C'est la feuille qui distingue un professionnel. On n'y écrit pas « j'ai utilisé l'IA » : "
                "on y écrit le prompt exact, la sortie, l'erreur trouvée, et le V du protocole qui l'a attrapée.")
    ws["A2"].font = Faide
    entete(ws, ["Étape", "Le prompt utilisé", "Ce que l'IA a rendu",
                "L'erreur détectée", "Le V qui l'a attrapée", "La correction"], ligne=4)
    for i in range(5, 10):
        for j in range(1, 7):
            ws.cell(i, j).border = Bord
    ws.cell(5, 1, "Lecture de la facture")
    ws.cell(6, 1, "Formule de recherche du taux")
    largeurs(ws, [24, 46, 34, 34, 20, 40])
    return ws

# ══════════════════════════════════════════════════════════════════════
#  5. Les classeurs des leçons
# ══════════════════════════════════════════════════════════════════════
def lecon_01():
    """Une plage morte. On la rend vivante."""
    for suffixe, propre in (("DEPART", False), ("CORRIGE", True)):
        wb = openpyxl.Workbook()
        ws = wb.active; ws.title = "Caisse"
        ecrire_caisse(ws, propre=propre, avec_montant=propre)
        if propre:
            tableau(ws, "t_Caisse", f"A1:I{len(LIGNES)+1}")
        wb.save(RACINE / f"M01_L01_{suffixe}.xlsx")

def lecon_02():
    """La table des taux arrive. Une seule cellule doit porter le taux."""
    for suffixe, corrige in (("DEPART", False), ("CORRIGE", True)):
        wb = openpyxl.Workbook()
        ws = wb.active; ws.title = "Caisse"
        ecrire_caisse(ws, propre=True, avec_montant=True)
        tableau(ws, "t_Caisse", f"A1:I{len(LIGNES)+1}")
        ecrire_taux(wb.create_sheet("TAUX"))
        s = wb.create_sheet("Synthese")
        s["A1"] = "Synthèse de septembre 2026"; s["A1"].font = Ftitre
        s["A3"] = "Mois analysé";        s["B3"] = MOIS_ANALYSE; s["B3"].number_format = "MMM YYYY"
        s["A4"] = "Taux CDF → USD";      s["A5"] = "CA total en CDF"
        s["A6"] = "CA total en USD";     s["A7"] = "Dernier PU — SUPERMARCHE CITY"
        if corrige:
            s["B4"] = '=_xlfn.XLOOKUP(B3,t_Taux[Mois],t_Taux[CDF_USD])'
            s["B5"] = '=SUM(t_Caisse[Montant])'
            s["B6"] = '=B5/TauxUSD'
            s["B7"] = ('=_xlfn.XLOOKUP("SUPERMARCHE CITY",t_Caisse[Client],'
                       't_Caisse[PU_CDF],"absent",0,-1)')
            wb.defined_names.add(DefinedName("TauxUSD", attr_text="Synthese!$B$4"))
        largeurs(s, [32, 20])
        wb.save(RACINE / f"M01_L02_{suffixe}.xlsx")

def lecon_03():
    """Trois formules, zéro clic, un tableau de bord qui se refait seul."""
    for suffixe, corrige in (("DEPART", False), ("CORRIGE", True)):
        wb = openpyxl.Workbook()
        ws = wb.active; ws.title = "Caisse"
        ecrire_caisse(ws, propre=True, avec_montant=True)
        tableau(ws, "t_Caisse", f"A1:I{len(LIGNES)+1}")
        b = wb.create_sheet("Bord")
        b["A1"] = "Le tableau de bord qui se fabrique tout seul"; b["A1"].font = Ftitre
        b["A3"] = "Seuil (CDF)"; b["B3"] = 500_000
        b["A5"] = "Clients distincts"
        b["A6"] = "Familles vendues"
        b["A7"] = "Ventes au-dessus du seuil"
        b["A9"] = "Les clients, triés A → Z"
        if corrige:
            b["B5"] = '=COUNTA(_xlfn.UNIQUE(t_Caisse[Client]))'
            b["B6"] = ('=_xlfn.TEXTJOIN("; ",TRUE,_xlfn._xlws.SORT('
                       '_xlfn.UNIQUE(t_Caisse[Famille])))')
            b["B7"] = '=ROWS(_xlfn._xlws.FILTER(t_Caisse[Ticket],t_Caisse[Montant]>B3))'
            b["A10"] = '=_xlfn._xlws.SORT(_xlfn.UNIQUE(t_Caisse[Client]))'
        largeurs(b, [32, 46])
        wb.save(RACINE / f"M01_L03_{suffixe}.xlsx")

def lecon_05():
    """La facture lue par l'IA, et le total de contrôle qui la juge."""
    for suffixe, corrige in (("DEPART", False), ("CORRIGE", True)):
        wb = openpyxl.Workbook()
        ws = wb.active; ws.title = "FACTURE"
        ws["A1"] = f"Facture fournisseur {NUM_FACTURE} — {DATE_FACTURE:%d/%m/%Y}"
        ws["A1"].font = Ftitre
        ws["A2"] = "Saisis ici ce que l'IA a lu sur la photo. Ne calcule rien à sa place."
        ws["A2"].font = Faide
        entete(ws, ["Reference", "Designation", "Quantite", "PU_CDF", "Montant"], ligne=4)
        n = len(FACTURE)
        for i in range(5, 5 + n):
            if corrige:
                r, d, q, p = FACTURE[i - 5]
                ws.cell(i, 1, r); ws.cell(i, 2, d); ws.cell(i, 3, q)
                ws.cell(i, 4, p); ws.cell(i, 5, q * p)
            for j in range(1, 6):
                ws.cell(i, j).border = Bord
        tableau(ws, "t_Facture", f"A4:E{4+n}")
        largeurs(ws, [14, 32, 12, 14, 16])

        c = wb.create_sheet("CONTROLE")
        c["A1"] = "Le contrôle en trois points"; c["A1"].font = Ftitre
        c["A3"] = "Total imprimé sur la facture"; c["B3"] = TOTAL_FACTURE
        c["A4"] = "Total recalculé depuis tes lignes"
        c["A5"] = "Écart — doit valoir 0"
        c["A6"] = "Nombre de lignes saisies"
        c["A7"] = "Nombre de lignes attendues";  c["B7"] = n
        if corrige:
            c["B4"] = "=SUM(t_Facture[Montant])"
            c["B5"] = "=B4-TotalImprime"
            c["B6"] = "=COUNTA(t_Facture[Reference])"
            wb.defined_names.add(DefinedName("TotalImprime", attr_text="CONTROLE!$B$3"))
        largeurs(c, [36, 20])
        wb.save(RACINE / f"M01_L05_{suffixe}.xlsx")

# ══════════════════════════════════════════════════════════════════════
#  6. LE TP 1 — un classeur à trous
#
#  Règle de conception : chaque réponse est UNE cellule, UNE formule,
#  UNE valeur scalaire. Jamais une matricielle nue — le correcteur ne
#  lirait que la première valeur du déversement.
# ══════════════════════════════════════════════════════════════════════
REPONSES = [
    (4,  "Nombre de mouvements enregistrés",
         '=COUNTA(t_Caisse[Ticket])'),
    (5,  "Chiffre d'affaires total, en CDF",
         '=SUM(t_Caisse[Montant])'),
    (6,  "Panier moyen par ticket, en CDF",
         '=AVERAGE(t_Caisse[Montant])'),
    (7,  "Taux CDF → USD du mois analysé",
         '=_xlfn.XLOOKUP(MoisAnalyse,t_Taux[Mois],t_Taux[CDF_USD])'),
    (8,  "Chiffre d'affaires total, en USD",
         '=B5/TauxUSD'),
    (9,  "Dernier prix unitaire appliqué à SUPERMARCHE CITY",
         '=_xlfn.XLOOKUP("SUPERMARCHE CITY",t_Caisse[Client],t_Caisse[PU_CDF],"absent",0,-1)'),
    (10, "Nombre de clients distincts",
         '=COUNTA(_xlfn.UNIQUE(t_Caisse[Client]))'),
    (11, "Familles de produits vendues, triées, séparées par « ; »",
         '=_xlfn.TEXTJOIN(";",TRUE,_xlfn._xlws.SORT(_xlfn.UNIQUE(t_Caisse[Famille])))'),
    (12, "Nombre de ventes supérieures à 500 000 CDF",
         '=ROWS(_xlfn._xlws.FILTER(t_Caisse[Ticket],t_Caisse[Montant]>500000))'),
    (13, "Quantité totale vendue (attention aux quantités en texte)",
         '=SUM(t_Caisse[Quantite])'),
    (14, "Total de la facture, recalculé depuis tes lignes",
         '=SUM(t_Facture[Montant])'),
    (15, "Écart facture : recalculé − imprimé  →  doit valoir 0",
         '=B14-TotalImprime'),
    (16, "Contrôle : lignes de Caisse − lignes de RAW  →  doit valoir 0",
         '=COUNTA(t_Caisse[Ticket])-COUNTA(RAW!B2:B61)'),
]

def tp01(corrige: bool):
    wb = openpyxl.Workbook()

    # RAW — intacte, avec ses six quantités en texte
    raw = wb.active; raw.title = "RAW"
    ecrire_caisse(raw, propre=False, avec_montant=False)

    # Caisse — c'est le travail
    caisse = wb.create_sheet("Caisse")
    if corrige:
        ecrire_caisse(caisse, propre=True, avec_montant=True)
        tableau(caisse, "t_Caisse", f"A1:I{len(LIGNES)+1}")
    else:
        entete(caisse, ["Date", "Ticket", "Client", "Produit", "Famille",
                        "Quantite", "PU_CDF", "Commercial", "Montant"])
        caisse["A3"] = ("Colle ici les 60 lignes de RAW, répare les quantités stockées "
                        "en texte, ajoute la colonne Montant, puis convertis en tableau "
                        "structuré nommé t_Caisse (Ctrl+L).")
        caisse["A3"].font = Faide
        largeurs(caisse, [12, 11, 26, 26, 12, 10, 14, 18, 16])

    ecrire_taux(wb.create_sheet("TAUX"))

    # FACTURE
    fac = wb.create_sheet("FACTURE")
    fac["A1"] = f"Facture fournisseur {NUM_FACTURE} — {DATE_FACTURE:%d/%m/%Y}"
    fac["A1"].font = Ftitre
    fac["A2"] = "Saisis ce que l'IA a lu sur la photo. Ne recalcule rien à sa place."
    fac["A2"].font = Faide
    entete(fac, ["Reference", "Designation", "Quantite", "PU_CDF", "Montant"], ligne=4)
    n = len(FACTURE)
    for i in range(5, 5 + n):
        if corrige:
            r, d, q, p = FACTURE[i - 5]
            fac.cell(i, 1, r); fac.cell(i, 2, d); fac.cell(i, 3, q)
            fac.cell(i, 4, p); fac.cell(i, 5, q * p)
        for j in range(1, 6):
            fac.cell(i, j).border = Bord
    tableau(fac, "t_Facture", f"A4:E{4+n}")
    largeurs(fac, [14, 32, 12, 14, 16])

    # REPONSES
    rep = wb.create_sheet("REPONSES")
    rep["A1"] = "REPONSES — une formule par cellule"; rep["A1"].font = Ftitre
    rep["A2"] = ("Une valeur tapée à la main ne rapporte que la moitié des points. "
                 "La machine regarde la formule autant que le résultat.")
    rep["A2"].font = Faide
    rep["A3"] = "Mois analysé"; rep["B3"] = MOIS_ANALYSE
    rep["B3"].number_format = "MMM YYYY"; rep["B3"].font = Fgras
    rep["A18"] = "Total imprimé sur la facture"; rep["B18"] = TOTAL_FACTURE
    rep["B18"].number_format = '# ##0 "CDF"'
    for ligne, libelle, formule in REPONSES:
        rep.cell(ligne, 1, libelle).border = Bord
        cel = rep.cell(ligne, 2)
        cel.border = Bord
        cel.fill = Rgris
        if corrige:
            cel.value = formule
    largeurs(rep, [56, 26])

    # Les trois plages nommées font partie du travail : les livrer toutes
    # faites offrirait trois points et supprimerait la leçon 1.2 — « une
    # seule cellule de taux, nommée, référencée par tout le classeur ».
    if corrige:
        wb.defined_names.add(DefinedName("MoisAnalyse",  attr_text="REPONSES!$B$3"))
        wb.defined_names.add(DefinedName("TauxUSD",      attr_text="REPONSES!$B$7"))
        wb.defined_names.add(DefinedName("TotalImprime", attr_text="REPONSES!$B$18"))
    else:
        rep["D3"]  = "← nomme cette cellule  MoisAnalyse"
        rep["D7"]  = "← nomme cette cellule  TauxUSD"
        rep["D18"] = "← nomme cette cellule  TotalImprime"
        for a in ("D3", "D7", "D18"):
            rep[a].font = Faide
        largeurs(rep, [56, 26, 4, 42])

    feuille_readme(wb, rempli=corrige)
    feuille_qualite(wb, len(LIGNES), None, rempli=corrige)
    feuille_annexe(wb)

    nom = "TP01_CORRIGE.xlsx" if corrige else "TP01_DEPART.xlsx"
    wb.save(RACINE / nom)
    return nom

# ══════════════════════════════════════════════════════════════════════
#  7. Les jeux de données bruts, livrés tels quels
# ══════════════════════════════════════════════════════════════════════
def jeux():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Caisse"
    ecrire_caisse(ws, propre=False, avec_montant=False)
    wb.save(RACINE / "J01_Caisse_Kinshasa.xlsx")

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Taux"
    ecrire_taux(ws)
    wb.save(RACINE / "J01_Taux_Change.xlsx")

if __name__ == "__main__":
    lecon_01(); lecon_02(); lecon_03(); lecon_05()
    jeux(); tp01(False); tp01(True)
    print("classeurs écrits dans", RACINE)
