#!/usr/bin/env python3
# ══════════════════════════════════════════════════════════════════════
#  Données du module 9 — décider
#
#  Un business plan : l'ouverture d'une septième agence à Bamako.
#  24 mois, 14 hypothèses, trois devises de fait (XOF encaissé,
#  USD décaissé, USD de reporting) et un taux de change qui bouge.
#
#  Tout est calculé ici, en Python, AVANT de fabriquer le moindre
#  classeur : ce sont ces valeurs-là qui serviront de référence pour
#  vérifier le corrigé recalculé par Excel.
# ══════════════════════════════════════════════════════════════════════
import datetime as dt, pathlib, sys, json, random

random.seed(909)
RACINE = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "sortie-m09")
RACINE.mkdir(parents=True, exist_ok=True)

# ══ 1 · Les quatorze hypothèses ═══════════════════════════════════════
#  L'ordre est celui de la feuille HYPOTHESES. Le code en colonne A
#  devient le nom de la plage : « Créer depuis la sélection » les pose
#  toutes les quatorze en deux clics.
H = [
 ("H_AMENAGEMENT",      68_000,    '#,##0\\ "USD"',
  "Aménagement du local — cloisons, rayonnage, groupe électrogène"),
 ("H_MATERIEL",         41_500,    '#,##0\\ "USD"',
  "Matériel roulant et informatique — une camionnette, six postes"),
 ("H_STOCK_INITIAL",    95_000,    '#,##0\\ "USD"',
  "Stock initial — payé en dollars, comme tous les achats du groupe"),
 ("H_VOLUME_M1",        4_200,     "#,##0",
  "Volume du premier mois, en unités"),
 ("H_VOLUME_CROISIERE", 19_500,    "#,##0",
  "Volume de croisière, en unités par mois"),
 ("H_MONTEE",           14,        "#,##0",
  "Durée de la montée en charge, en mois — calibrée sur Libreville"),
 ("H_PRIX_XOF",         4_350,     '#,##0\\ "XOF"',
  "Prix de vente moyen par unité — encaissé en francs CFA"),
 ("H_COUT_USD",         4.55,      '#,##0.00\\ "USD"',
  "Coût d'achat moyen par unité — décaissé en dollars"),
 ("H_LOYER_XOF",        2_750_000, '#,##0\\ "XOF"',
  "Loyer mensuel du local"),
 ("H_PERSONNEL_XOF",    6_480_000, '#,##0\\ "XOF"',
  "Masse salariale mensuelle — onze personnes"),
 ("H_AUTRES_XOF",       1_920_000, '#,##0\\ "XOF"',
  "Autres charges mensuelles — énergie, transport, assurances"),
 ("H_TAUX_XOF",         631.40,    '#,##0.00',
  "Taux de change au démarrage, en XOF pour 1 USD"),
 ("H_DERIVE_TAUX",      0.0045,    "0.00%",
  "Dérive mensuelle du taux retenue au plan — soit 5,5 % par an"),
 ("H_ACTUALISATION",    0.14,      "0.00%",
  "Taux d'actualisation annuel exigé par la direction"),
]
HV = {c: v for c, v, _f, _d in H}

SCENARIOS = [
 ("Pessimiste", 16_500, 4_200, 0.0080),
 ("Base",       19_500, 4_350, 0.0045),
 ("Optimiste",  22_000, 4_450, 0.0020),
]

# ══ 2 · Le modèle, en Python ══════════════════════════════════════════
def projection(vc, prix, derive):
    """Les 24 flux mensuels en USD, et ce qu'il faut pour les expliquer.

    On vend en francs CFA, on achète en dollars : c'est cette asymétrie
    qui met le taux de change au centre du plan, et pas en annexe."""
    v1, montee = HV["H_VOLUME_M1"], HV["H_MONTEE"]
    cout, T0 = HV["H_COUT_USD"], HV["H_TAUX_XOF"]
    charges = HV["H_LOYER_XOF"] + HV["H_PERSONNEL_XOF"] + HV["H_AUTRES_XOF"]
    I0 = HV["H_AMENAGEMENT"] + HV["H_MATERIEL"] + HV["H_STOCK_INITIAL"]
    i = (1 + HV["H_ACTUALISATION"]) ** (1 / 12) - 1

    lignes = [dict(mois=0, taux=T0, volume=0.0, ca_xof=0.0, charges_xof=0.0,
                   ca_usd=0.0, charges_usd=0.0, marge_usd=0.0, achats=0.0,
                   flux=-I0, actu=1.0, flux_act=-I0, cumul=-I0)]
    for m in range(1, 25):
        v = min(vc, v1 + (vc - v1) * (m - 1) / (montee - 1))
        tx = T0 * (1 + derive) ** m
        ca_xof = v * prix
        ca_usd, ch_usd = ca_xof / tx, charges / tx
        achats = v * cout
        flux = (ca_usd - ch_usd) - achats
        actu = 1 / (1 + i) ** m
        lignes.append(dict(mois=m, taux=tx, volume=v, ca_xof=ca_xof,
                           charges_xof=charges, ca_usd=ca_usd,
                           charges_usd=ch_usd, marge_usd=ca_usd - ch_usd,
                           achats=achats, flux=flux, actu=actu,
                           flux_act=flux * actu,
                           cumul=lignes[-1]["cumul"] + flux))
    return lignes, i, I0, charges

def van(lignes, i, I0):
    return -I0 + sum(l["flux"] * l["actu"] for l in lignes[1:])

def tri_mensuel(lignes):
    """Bissection : plus lente que Newton, mais elle ne diverge jamais.
    C'est exactement la différence avec TRI, qui, elle, peut échouer."""
    flux = [l["flux"] for l in lignes]
    lo, hi = -0.95, 3.0
    for _ in range(400):
        mid = (lo + hi) / 2
        v = sum(f / (1 + mid) ** k for k, f in enumerate(flux))
        if v > 0: lo = mid
        else:     hi = mid
    return (lo + hi) / 2

BASE, i_mens, I0, CHARGES = projection(HV["H_VOLUME_CROISIERE"],
                                       HV["H_PRIX_XOF"], HV["H_DERIVE_TAUX"])
PESS = projection(*SCENARIOS[0][1:])[0]
OPTI = projection(*SCENARIOS[2][1:])[0]

VAN_BASE = van(BASE, i_mens, I0)
VAN_PESS = van(PESS, i_mens, I0)
VAN_OPTI = van(OPTI, i_mens, I0)
TRI_ANNUEL = (1 + tri_mensuel(BASE)) ** 12 - 1

#  Le taux limite. VAN = P/T0 - Q, avec P la valeur actuelle de ce qui
#  rentre en francs CFA et Q celle de ce qui sort en dollars. La VAN
#  s'annule donc en T0* = T0 x (entrées actualisées / sorties actualisées).
ENTREES = sum(l["marge_usd"] * l["actu"] for l in BASE[1:])
SORTIES = I0 + sum(l["achats"] * l["actu"] for l in BASE[1:])
TAUX_LIMITE = HV["H_TAUX_XOF"] * ENTREES / SORTIES

SEUIL = HV["H_LOYER_XOF"] + HV["H_PERSONNEL_XOF"] + HV["H_AUTRES_XOF"]
SEUIL = SEUIL / (HV["H_PRIX_XOF"] - HV["H_COUT_USD"] * HV["H_TAUX_XOF"])
import math
SEUIL_UNITES = math.ceil(SEUIL)
RETOUR = sum(1 for l in BASE if l["cumul"] < 0)

# ══ 3 · La table de données à deux entrées ════════════════════════════
#  20 volumes x 21 taux = 420 simulations. C'est ce rectangle qu'Excel
#  remplit sans macro, et que presque personne n'a jamais utilisé.
VOLUMES = [15_000 + 500 * k for k in range(20)]
TAUX_AXE = [600 + 5 * k for k in range(21)]

def van_pour(taux0, vc):
    v1, montee = HV["H_VOLUME_M1"], HV["H_MONTEE"]
    cout, prix, derive = HV["H_COUT_USD"], HV["H_PRIX_XOF"], HV["H_DERIVE_TAUX"]
    total = -I0
    for m in range(1, 25):
        v = min(vc, v1 + (vc - v1) * (m - 1) / (montee - 1))
        tx = taux0 * (1 + derive) ** m
        flux = (v * prix - CHARGES) / tx - v * cout
        total += flux / (1 + i_mens) ** m
    return total

TABLE2 = {(t, v): van_pour(t, v) for t in TAUX_AXE for v in VOLUMES}
VAN_660_18000 = TABLE2[(660, 18_000)]

# ══ 4 · L'historique de Libreville — ce qui calibre la montée ═════════
#  L'hypothèse de 14 mois n'est pas un chiffre rond sorti d'une réunion :
#  c'est la dernière ouverture du groupe, relue.
MOIS_FR = ["janv", "févr", "mars", "avr", "mai", "juin",
           "juil", "août", "sept", "oct", "nov", "déc"]
LIBREVILLE = []
_croisiere = 17_800
for m in range(1, 25):
    part = min(1.0, 0.24 + 0.76 * (m - 1) / 12.5)
    v = round(_croisiere * part * random.uniform(0.94, 1.06))
    d = dt.date(2024 + (m - 1) // 12, (m - 1) % 12 + 1, 1)
    LIBREVILLE.append({
        "periode": f"{MOIS_FR[d.month - 1]}-{str(d.year)[2:]}",
        "date": d, "mois": m, "volume": v,
        "ca_xaf": round(v * random.uniform(3_980, 4_310)),
    })
VOL_CROISIERE_LBV = max(l["volume"] for l in LIBREVILLE)
#  Le mois où Libreville a atteint 95 % de son volume de croisière.
MOIS_95 = next(l["mois"] for l in LIBREVILLE
               if l["volume"] >= 0.95 * VOL_CROISIERE_LBV)

R = {
 "investissement": I0,
 "charges_mensuelles_xof": CHARGES,
 "taux_mensuel": round(i_mens, 10),
 "taux_mois_24": round(HV["H_TAUX_XOF"] * (1 + HV["H_DERIVE_TAUX"]) ** 24, 6),
 "volume_total": round(sum(l["volume"] for l in BASE), 2),
 "volume_mois_12": round(BASE[12]["volume"], 4),
 "ca_cumule_usd": round(sum(l["ca_usd"] for l in BASE), 2),
 "achats_cumules_usd": round(sum(l["achats"] for l in BASE), 2),
 "charges_cumulees_usd": round(sum(l["charges_usd"] for l in BASE), 2),
 "marge_brute_usd": round(sum(l["ca_usd"] for l in BASE)
                          - sum(l["achats"] for l in BASE), 2),
 "resultat_cumule_usd": round(sum(l["flux"] for l in BASE[1:]), 2),
 "van": round(VAN_BASE, 2),
 "tri_annuel": round(TRI_ANNUEL, 6),
 "seuil_unites": SEUIL_UNITES,
 "mois_croisiere": HV["H_MONTEE"],
 "mois_retour": RETOUR,
 "meilleur_mois_usd": round(max(l["flux"] for l in BASE[1:]), 2),
 "van_pessimiste": round(VAN_PESS, 2),
 "van_optimiste": round(VAN_OPTI, 2),
 "taux_limite": round(TAUX_LIMITE, 4),
 "marge_securite_taux": round(TAUX_LIMITE / HV["H_TAUX_XOF"] - 1, 6),
 "van_660_18000": round(VAN_660_18000, 2),
 "libreville_croisiere": VOL_CROISIERE_LBV,
 "libreville_mois_95": MOIS_95,
 "simulations": len(TABLE2),
}
for k, v in R.items():
    print(f"  {k:<26} {v}")

# ══ 5 · Les classeurs ═════════════════════════════════════════════════
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

Ftitre  = Font(name="Calibri", size=14, bold=True)
Fentete = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
Faide   = Font(name="Calibri", size=10, italic=True, color="FF6B7280")
Fgras   = Font(name="Calibri", size=11, bold=True)
#  La convention des modélisateurs financiers. Elle n'est pas
#  décorative : c'est un contrôle visuel, et c'est tout le module 9.3.
Fbleu   = Font(name="Calibri", size=11, color="FF0000FF")          # saisie
Fnoir   = Font(name="Calibri", size=11, color="FF000000")          # formule
Fvert   = Font(name="Calibri", size=11, color="FF008000")          # lien interne
Rentete = PatternFill("solid", fgColor="FF243044")
Rjaune  = PatternFill("solid", fgColor="FFFFF6D8")
Rbleu   = PatternFill("solid", fgColor="FFEFF4FE")
Rvert   = PatternFill("solid", fgColor="FFE8F6EE")
Bord    = Border(bottom=Side(style="thin", color="FFD6DBE1"))

FMT_USD = '#,##0.00\\ "USD"'
FMT_USD0 = '#,##0\\ "USD"'
FMT_XOF = '#,##0\\ "XOF"'
FMT_PCT = "0.00%"
FMT_INT = "#,##0"
FMT_TAUX = "#,##0.00"

def entete(ws, cols, ligne=1, depart=1):
    for k, c in enumerate(cols, start=depart):
        cel = ws.cell(ligne, k, c)
        cel.font, cel.fill = Fentete, Rentete
        cel.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[ligne].height = 22

def colonnes(ws, specs, depart=1):
    for k, (w, fmt) in enumerate(specs, start=depart):
        d = ws.column_dimensions[get_column_letter(k)]
        d.width = w
        if fmt: d.number_format = fmt

def tableau(ws, nom, ref):
    t = Table(displayName=nom, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)

# ── HYPOTHESES ────────────────────────────────────────────────────────
def feuille_hypotheses(wb):
    ws = wb.create_sheet("HYPOTHESES")
    colonnes(ws, [(26, None), (18, None), (72, None)])
    ws["A1"] = "Les quatorze hypothèses"; ws["A1"].font = Ftitre
    ws["A2"] = ("Tout ce que ce plan suppose est ici, et nulle part ailleurs. "
                "Une cellule bleue hors de cette feuille est un bug.")
    ws["A2"].font = Faide
    entete(ws, ["Code", "Valeur", "Ce que c'est"], ligne=4)
    for k, (code, val, fmt, quoi) in enumerate(H, start=5):
        c = ws.cell(k, 1, code); c.font, c.border = Fgras, Bord
        d = ws.cell(k, 2, val)
        d.font, d.fill, d.border, d.number_format = Fbleu, Rbleu, Bord, fmt
        ws.cell(k, 3, quoi).font = Faide

    ws["A20"] = "Les trois scénarios"; ws["A20"].font = Ftitre
    entete(ws, ["Scenario", "Volume_croisiere", "Prix_XOF", "Derive_Taux"], ligne=21)
    for k, (nom, vc, prix, d) in enumerate(SCENARIOS, start=22):
        ws.cell(k, 1, nom).font = Fgras
        for col, (v, fmt) in enumerate(((vc, FMT_INT), (prix, FMT_XOF),
                                        (d, FMT_PCT)), start=2):
            c = ws.cell(k, col, v); c.font, c.fill, c.number_format = Fbleu, Rbleu, fmt
    tableau(ws, "t_Scenarios", "A21:D24")
    ws["A26"] = ("La ligne « Base » doit être identique aux hypothèses du haut. "
                 "La feuille QUALITE le vérifie.")
    ws["A26"].font = Faide
    return ws

# ── RAW · l'ouverture de Libreville, telle qu'elle est sortie du SI ────
def feuille_raw(wb):
    ws = wb.create_sheet("RAW")
    colonnes(ws, [(14, None), (10, None), (14, None), (16, None)])
    ws["A1"] = "Libreville — les 24 premiers mois, extraction brute"
    ws["A1"].font = Ftitre
    ws["A2"] = ("Source : export du SI commercial, 3 septembre 2026. "
                "On n'y touche pas — on la copie vers CLEAN.")
    ws["A2"].font = Faide
    entete(ws, ["Periode", "Mois", "Volume", "CA_XAF"], ligne=4)
    for k, l in enumerate(LIBREVILLE, start=5):
        ws.cell(k, 1, l["periode"]).border = Bord
        ws.cell(k, 2, l["mois"]).border = Bord
        ws.cell(k, 3, l["volume"]).border = Bord
        c = ws.cell(k, 4, l["ca_xaf"]); c.border, c.number_format = Bord, "#,##0"
    tableau(ws, "t_Libreville", f"A4:D{4 + len(LIBREVILLE)}")
    return ws

# ── CLEAN · ce qui calibre la montée en charge ────────────────────────
def feuille_clean(wb, corrige):
    ws = wb.create_sheet("CLEAN")
    colonnes(ws, [(10, None), (14, None), (18, None), (46, None)])
    ws["A1"] = "Libreville — la courbe de montée en charge"; ws["A1"].font = Ftitre
    ws["A2"] = ("C'est d'ici que sort H_MONTEE. L'hypothèse n'est pas un chiffre "
                "rond : c'est la dernière ouverture du groupe, relue.")
    ws["A2"].font = Faide
    entete(ws, ["Mois", "Volume", "Part_croisiere", "Lecture"], ligne=4)
    for k, l in enumerate(LIBREVILLE, start=5):
        ws.cell(k, 1, l["mois"]).border = Bord
        ws.cell(k, 2, l["volume"]).border = Bord
        c = ws.cell(k, 3)
        c.border, c.number_format = Bord, FMT_PCT
        if corrige:
            c.value = f"=B{k}/MAX($B$5:$B$28)"
            c.font = Fnoir
    ws.cell(30, 1, "Volume de croisière observé").font = Fgras
    ws.cell(30, 2, VOL_CROISIERE_LBV).number_format = FMT_INT
    ws.cell(31, 1, "Mois où 95 % est atteint").font = Fgras
    ws.cell(31, 2, MOIS_95)
    ws.cell(32, 1, "Hypothèse retenue pour Bamako").font = Fgras
    ws.cell(32, 2, HV["H_MONTEE"])
    ws.cell(32, 4, "Un mois de plus que Libreville : le marché malien est "
                   "moins dense.").font = Faide
    return ws

# ── CALCULS · la projection sur 24 mois ───────────────────────────────
#  Rangée 5 : les en-têtes. Rangées 6 à 30 : les mois 0 à 24.
#  La disposition appartient à l'apprenant — la plage nommée LIBRE_CALCULS
#  dit au correcteur de ne pas noter ces cellules une à une.
CAL_COLS = [
 (9,  "Mois",              FMT_INT),
 (13, "Taux_XOF_USD",      FMT_TAUX),
 (12, "Volume",            FMT_INT),
 (16, "CA_XOF",            FMT_XOF),
 (16, "Charges_XOF",       FMT_XOF),
 (14, "CA_USD",            FMT_USD),
 (14, "Charges_USD",       FMT_USD),
 (14, "Marge_locale_USD",  FMT_USD),
 (14, "Achats_USD",        FMT_USD),
 (14, "Flux_USD",          FMT_USD),
 (14, "Actualisation",     "0.000000"),
 (14, "Flux_actualise",    FMT_USD),
 (15, "Cumul_USD",         FMT_USD),
]

def feuille_calculs(wb, corrige):
    ws = wb.create_sheet("CALCULS")
    colonnes(ws, [(w, None) for w, _n, _f in CAL_COLS])
    ws["A1"] = "La projection sur 24 mois"; ws["A1"].font = Ftitre
    ws["A2"] = ("Une ligne par mois, le mois 0 portant l'investissement. "
                "Aucune valeur en dur : tout descend de HYPOTHESES.")
    ws["A2"].font = Faide
    ws["A3"] = "Taux d'actualisation mensuel"; ws["A3"].font = Fgras
    ws["B3"].number_format = "0.0000000"
    entete(ws, [n for _w, n, _f in CAL_COLS], ligne=5)
    for r in range(6, 31):
        for k, (_w, _n, fmt) in enumerate(CAL_COLS, start=1):
            c = ws.cell(r, k); c.border, c.number_format = Bord, fmt
    if not corrige:
        return ws

    ws["B3"] = "=(1+H_ACTUALISATION)^(1/12)-1"; ws["B3"].font = Fnoir
    F = {}
    for r in range(6, 31):
        m = r - 6
        F[f"A{r}"] = "=0" if m == 0 else f"=A{r-1}+1"
        F[f"B{r}"] = f"=H_TAUX_XOF*(1+H_DERIVE_TAUX)^A{r}"
        F[f"C{r}"] = ("=0" if m == 0 else
                      "=MIN(H_VOLUME_CROISIERE,H_VOLUME_M1+"
                      f"(H_VOLUME_CROISIERE-H_VOLUME_M1)*(A{r}-1)/(H_MONTEE-1))")
        F[f"D{r}"] = f"=C{r}*H_PRIX_XOF"
        F[f"E{r}"] = ("=0" if m == 0 else
                      "=H_LOYER_XOF+H_PERSONNEL_XOF+H_AUTRES_XOF")
        F[f"F{r}"] = f"=D{r}/B{r}"
        F[f"G{r}"] = f"=E{r}/B{r}"
        F[f"H{r}"] = f"=F{r}-G{r}"
        F[f"I{r}"] = f"=C{r}*H_COUT_USD"
        F[f"J{r}"] = ("=-(H_AMENAGEMENT+H_MATERIEL+H_STOCK_INITIAL)" if m == 0
                      else f"=H{r}-I{r}")
        F[f"K{r}"] = f"=1/(1+$B$3)^A{r}"
        F[f"L{r}"] = f"=J{r}*K{r}"
        F[f"M{r}"] = f"=J{r}" if m == 0 else f"=M{r-1}+J{r}"
    for ref, f in F.items():
        ws[ref] = f; ws[ref].font = Fnoir
    return ws

# ── ANALYSE · les scénarios recalculés, et les 420 simulations ────────
def feuille_analyse(wb, corrige):
    ws = wb.create_sheet("ANALYSE")
    colonnes(ws, [(10, None)] + [(16, None)] * 3)
    ws["A1"] = "Les trois scénarios, recalculés côte à côte"; ws["A1"].font = Ftitre
    ws["A2"] = ("Le Gestionnaire de scénarios explore. Ce bloc-ci documente : "
                "il reste dans le classeur, et il se recalcule.")
    ws["A2"].font = Faide
    entete(ws, ["Mois", "Pessimiste", "Base", "Optimiste"], ligne=5)
    for r in range(6, 31):
        for k in range(1, 5):
            c = ws.cell(r, k); c.border
            c.number_format = FMT_INT if k == 1 else FMT_USD

    ws["A33"] = "Table de données à deux entrées — VAN en USD"
    ws["A33"].font = Ftitre
    ws["A34"] = (f"{len(TABLE2)} simulations. En ligne les volumes de croisière, "
                 "en colonne le taux de change de départ.")
    ws["A34"].font = Faide
    if not corrige:
        return ws

    for r in range(6, 31):
        m = r - 6
        ws.cell(r, 1, "=CALCULS!A%d" % r).font = Fvert
        for k in range(2, 5):
            c = ws.cell(r, k)
            if m == 0:
                c.value = "=-(H_AMENAGEMENT+H_MATERIEL+H_STOCK_INITIAL)"
            else:
                j = k - 1
                vc = f"INDEX(t_Scenarios[Volume_croisiere],{j})"
                px = f"INDEX(t_Scenarios[Prix_XOF],{j})"
                dd = f"INDEX(t_Scenarios[Derive_Taux],{j})"
                v = (f"MIN({vc},H_VOLUME_M1+({vc}-H_VOLUME_M1)"
                     f"*($A{r}-1)/(H_MONTEE-1))")
                tx = f"H_TAUX_XOF*(1+{dd})^$A{r}"
                c.value = (f"=({v}*{px}-(H_LOYER_XOF+H_PERSONNEL_XOF"
                           f"+H_AUTRES_XOF))/({tx})-{v}*H_COUT_USD")
            c.font = Fnoir

    #  La table de données est écrite en valeurs : openpyxl ne sait pas
    #  poser une vraie table Excel, et c'est justement le geste que le TP
    #  demande. Le rectangle sert de référence de lecture.
    ws.cell(36, 1, "VAN").font = Fgras
    for k, v in enumerate(VOLUMES, start=2):
        c = ws.cell(36, k, v); c.font, c.number_format = Fgras, FMT_INT
        ws.column_dimensions[get_column_letter(k)].width = 12
    for r, t in enumerate(TAUX_AXE, start=37):
        c = ws.cell(r, 1, t); c.font, c.number_format = Fgras, FMT_TAUX
        for k, v in enumerate(VOLUMES, start=2):
            d = ws.cell(r, k, round(TABLE2[(t, v)], 2))
            d.number_format = "#,##0"
    return ws

# ── REPONSES ──────────────────────────────────────────────────────────
#  ① ce qu'on lit dans les outils — tapé, donc non noté par la machine.
LUS = [
 ("Taux limite trouvé par Valeur cible, en XOF pour 1 USD",
  FMT_TAUX, round(TAUX_LIMITE, 2)),
 ("VAN du scénario pessimiste, lue dans le rapport de synthèse",
  FMT_USD0, round(VAN_PESS, 2)),
 ("VAN au croisement taux 660 / volume 18 000 de la table de données",
  FMT_USD0, round(VAN_660_18000, 2)),
]

#  ② ce que les formules vérifient — noté. Rangées 11 à 36.
REPONSES = [
 ("Investissement initial total", "=H_AMENAGEMENT+H_MATERIEL+H_STOCK_INITIAL",
  FMT_USD0, R["investissement"]),
 ("Charges fixes mensuelles", "=H_LOYER_XOF+H_PERSONNEL_XOF+H_AUTRES_XOF",
  FMT_XOF, R["charges_mensuelles_xof"]),
 ("Taux d'actualisation mensuel", "=(1+H_ACTUALISATION)^(1/12)-1",
  "0.0000%", R["taux_mensuel"]),
 ("Taux de change au mois 24", "=H_TAUX_XOF*(1+H_DERIVE_TAUX)^24",
  FMT_TAUX, R["taux_mois_24"]),
 ("Volume total vendu sur 24 mois", "=SUM(F_VOLUME)", FMT_INT, R["volume_total"]),
 ("Volume du mois 12", "=INDEX(F_VOLUME,12)", FMT_INT, R["volume_mois_12"]),
 ("Chiffre d'affaires cumulé, en USD", "=SUM(F_CA_USD)", FMT_USD,
  R["ca_cumule_usd"]),
 ("Achats cumulés, en USD", "=SUM(F_ACHATS)", FMT_USD, R["achats_cumules_usd"]),
 ("Charges cumulées, en USD", "=SUM(F_CHARGES_USD)", FMT_USD,
  R["charges_cumulees_usd"]),
 ("Marge brute cumulée, en USD", "=B17-B18", FMT_USD, R["marge_brute_usd"]),
 ("Taux de marge brute moyen", "=1-B18/B17", FMT_PCT, None),
 ("Résultat cumulé sur 24 mois, en USD", "=SUM(F_FLUX1)", FMT_USD,
  R["resultat_cumule_usd"]),
 ("VAN du projet, en USD  ·  le flux du mois 0 n'entre pas dans VAN",
  "=NPV(B13,F_FLUX1)-B11", FMT_USD, R["van"]),
 ("TRI annuel du projet", "=(1+IRR(F_FLUX))^12-1", FMT_PCT, R["tri_annuel"]),
 ("Seuil de rentabilité mensuel, en unités",
  "=ROUNDUP(B12/(H_PRIX_XOF-H_COUT_USD*H_TAUX_XOF),0)", FMT_INT,
  R["seuil_unites"]),
 ("Mois où le volume de croisière est atteint",
  '=COUNTIF(F_VOLUME,"<"&H_VOLUME_CROISIERE)+1', FMT_INT, R["mois_croisiere"]),
 ("Mois de retour sur investissement", '=COUNTIF(F_CUMUL,"<0")', FMT_INT,
  R["mois_retour"]),
 ("Meilleur mois en résultat, en USD", "=MAX(F_FLUX1)", FMT_USD,
  R["meilleur_mois_usd"]),
 ("VAN du scénario pessimiste", "=NPV(B13,S_FLUX_PESS)-B11", FMT_USD,
  R["van_pessimiste"]),
 ("VAN du scénario optimiste", "=NPV(B13,S_FLUX_OPT)-B11", FMT_USD,
  R["van_optimiste"]),
 ("Écart entre optimiste et pessimiste", "=B30-B29", FMT_USD, None),
 ("💎  Taux de change au-delà duquel le projet devient perdant",
  "=H_TAUX_XOF*SUMPRODUCT(F_MARGE_USD,F_ACTU)/(B11+SUMPRODUCT(F_ACHATS,F_ACTU))",
  FMT_TAUX, R["taux_limite"]),
 ("Marge de sécurité sur le taux de change", "=B32/H_TAUX_XOF-1", FMT_PCT,
  R["marge_securite_taux"]),
 ("Valeurs en dur hors HYPOTHESES  ·  doit valoir 0",
  '=SUMPRODUCT((Z_AUDIT<>"")*NOT(_xlfn.ISFORMULA(Z_AUDIT)))', FMT_INT, 0),
 ("R_ECART_TAUX — Valeur cible moins formule  ·  doit valoir 0",
  "=ROUND(B5-B32,0)", FMT_TAUX, 0),
 ("R_CONTROLE — résultat moins (CA − achats − charges)  ·  doit valoir 0",
  "=ROUND(B22-(B17-B18-B19),2)", FMT_USD, 0),
]
R["taux_marge_brute"] = round(1 - R["achats_cumules_usd"] / R["ca_cumule_usd"], 6)
R["ecart_scenarios"] = round(R["van_optimiste"] - R["van_pessimiste"], 2)
REPONSES[10] = REPONSES[10][:3] + (R["taux_marge_brute"],)
REPONSES[20] = REPONSES[20][:3] + (R["ecart_scenarios"],)

def feuille_reponses(wb, corrige):
    ws = wb.create_sheet("REPONSES")
    colonnes(ws, [(70, None), (22, None)])
    ws["A1"] = "TP 9 — La septième agence"; ws["A1"].font = Ftitre
    ws["A2"] = ("Les cellules jaunes se remplissent une fois le modèle construit. "
                "Les trois premières se lisent dans les outils ; les autres se calculent.")
    ws["A2"].font = Faide
    ws["A4"] = "①  CE QUE VOUS LISEZ DANS LES OUTILS"; ws["A4"].font = Fgras
    for k, (lib, fmt, val) in enumerate(LUS, start=5):
        ws.cell(k, 1, lib).border = Bord
        c = ws.cell(k, 2); c.fill, c.border, c.number_format = Rjaune, Bord, fmt
        if corrige: c.value, c.font = val, Fbleu
    ws["A9"] = "②  CE QUE LES FORMULES VÉRIFIENT"; ws["A9"].font = Fgras
    for k, (lib, f, fmt, _v) in enumerate(REPONSES, start=11):
        ws.cell(k, 1, lib).border = Bord
        c = ws.cell(k, 2); c.fill, c.border, c.number_format = Rjaune, Bord, fmt
        if corrige: c.value, c.font = f, Fnoir
    return ws

# ── QUALITE ───────────────────────────────────────────────────────────
QUALITE = [
 ("Hypothèses renseignées  ·  doit valoir 14",
  "=COUNT(HYPOTHESES!$B$5:$B$18)", FMT_INT, 14,
  "quatorze, ni plus ni moins : une hypothèse de plus est une hypothèse cachée"),
 ("Hypothèses manquantes  ·  doit valoir 0",
  "=14-COUNT(HYPOTHESES!$B$5:$B$18)", FMT_INT, 0,
  "le même contrôle, écrit dans le sens où on le lit"),
 ("Volume du mois 24 moins la croisière  ·  doit valoir 0",
  "=ROUND(INDEX(F_VOLUME,24)-H_VOLUME_CROISIERE,0)", FMT_INT, 0,
  "si la montée en charge n'atteint jamais la croisière, la courbe est fausse"),
 ("Cumul du mois 24 moins la somme des flux  ·  doit valoir 0",
  "=ROUND(INDEX(F_CUMUL,25)-(SUM(F_FLUX1)-REPONSES!$B$11),2)", FMT_USD, 0,
  "le cumul et la somme doivent tomber d'accord, sinon une ligne a glissé"),
 ("Scénario Base moins l'hypothèse de volume  ·  doit valoir 0",
  "=ROUND(INDEX(t_Scenarios[Volume_croisiere],2)-H_VOLUME_CROISIERE,0)",
  FMT_INT, 0, "le scénario de base EST le plan : deux chiffres, une vérité"),
 ("Scénario Base moins l'hypothèse de prix  ·  doit valoir 0",
  "=ROUND(INDEX(t_Scenarios[Prix_XOF],2)-H_PRIX_XOF,0)", FMT_INT, 0,
  "le piège classique : on modifie une hypothèse et on oublie le scénario"),
 ("La VAN est-elle positive ?",
  '=IF(REPONSES!$B$23>0,"PASS","FAIL")', None, "PASS",
  "PASS ne veut pas dire « il faut y aller » : lisez la marge de sécurité"),
 ("Les trois scénarios sont-ils ordonnés ?",
  '=IF(AND(REPONSES!$B$29<REPONSES!$B$23,REPONSES!$B$23<REPONSES!$B$30),'
  '"PASS","FAIL")', None, "PASS",
  "pessimiste < base < optimiste : sinon un paramètre est inversé"),
]

def feuille_qualite(wb, corrige):
    ws = wb.create_sheet("QUALITE")
    colonnes(ws, [(56, None), (18, None), (68, None)])
    ws["A1"] = "Contrôles de qualité"; ws["A1"].font = Ftitre
    ws["A2"] = ("Huit contrôles. Tant qu'un seul ne passe pas, le modèle "
                "n'est pas prêt à sortir.")
    ws["A2"].font = Faide
    entete(ws, ["Contrôle", "Valeur", "Ce que ça veut dire"], ligne=4)
    for k, (lib, f, fmt, _v, quoi) in enumerate(QUALITE, start=5):
        ws.cell(k, 1, lib).border = Bord
        c = ws.cell(k, 2); c.fill, c.border = Rjaune, Bord
        if fmt: c.number_format = fmt
        if corrige: c.value, c.font = f, Fnoir
        ws.cell(k, 3, quoi).font = Faide
    return ws

# ── DASHBOARD ─────────────────────────────────────────────────────────
TABLEAU_BORD = [
 ("Investissement",            "=REPONSES!$B$11", FMT_USD0),
 ("VAN à 14 %",                "=REPONSES!$B$23", FMT_USD0),
 ("TRI annuel",                "=REPONSES!$B$24", FMT_PCT),
 ("Retour sur investissement", "=REPONSES!$B$27", '#,##0" mois"'),
 ("Seuil mensuel",             "=REPONSES!$B$25", '#,##0" unités"'),
 ("Taux limite",               "=REPONSES!$B$32", FMT_TAUX),
 ("Marge de sécurité",         "=REPONSES!$B$33", FMT_PCT),
 ("VAN pessimiste",            "=REPONSES!$B$29", FMT_USD0),
 ("VAN optimiste",             "=REPONSES!$B$30", FMT_USD0),
]

def feuille_dashboard(wb, corrige):
    ws = wb.create_sheet("DASHBOARD")
    colonnes(ws, [(34, None), (22, None)])
    ws["A1"] = "Bamako — la décision en une page"; ws["A1"].font = Ftitre
    ws["A2"] = ("Neuf chiffres. Si la direction en veut un dixième, c'est que "
                "l'un des neuf ne répond pas.")
    ws["A2"].font = Faide
    for k, (lib, f, fmt) in enumerate(TABLEAU_BORD, start=4):
        ws.cell(k, 1, lib).font = Fgras
        c = ws.cell(k, 2); c.number_format, c.border = fmt, Bord
        if corrige: c.value, c.font = f, Fvert
    ws.cell(15, 1, "La phrase à dire en réunion").font = Fgras
    if corrige:
        ws.cell(16, 1,
                "Le projet est rentable au plan de base, et il cesse de l'être "
                f"si le franc CFA part de {R['taux_limite']:.0f} au lieu de "
                f"{HV['H_TAUX_XOF']:.2f} — soit "
                f"{R['marge_securite_taux']*100:.1f} % de marge. "
                "Le reste se discute ; ça, non.").alignment = Alignment(wrap_text=True)
        ws.merge_cells("A16:B19")
    return ws

# ── README · ANNEXE_IA ────────────────────────────────────────────────
TEXTES = {
 "README": [
  ("À quoi sert ce classeur",
   "Décider de l'ouverture d'une septième agence BAOBAB à Bamako (Mali). "
   "Horizon 24 mois, du 1er janvier 2027 au 31 décembre 2028."),
  ("Qui l'a produit, et quand",
   "Direction financière BAOBAB, septembre 2026. Version 1.0."),
  ("Les sources",
   "Les investissements viennent des trois devis reçus en août 2026. "
   "La montée en charge est calibrée sur l'ouverture de Libreville "
   f"(feuilles RAW et CLEAN) : 95 % de la croisière atteints au mois "
   f"{MOIS_95}, retenu à {HV['H_MONTEE']} mois pour Bamako, le marché "
   "malien étant moins dense."),
  ("La règle des couleurs",
   "Bleu = une saisie, et toutes les saisies sont sur HYPOTHESES. "
   "Noir = une formule. Vert = un lien vers une autre feuille du classeur. "
   "Rouge = un lien vers un autre classeur — il n'y en a aucun ici, "
   "et c'est voulu."),
  ("Pourquoi le taux de change est au centre",
   "On encaisse en francs CFA et on décaisse en dollars : le stock est "
   "importé et payé par le groupe. Le résultat en dollars dépend donc "
   "directement du taux. Et le franc CFA n'est pas une devise flottante : "
   "il est arrimé à l'euro à parité fixe (655,957 XOF pour 1 EUR). "
   "Le « risque de change » de ce plan est en réalité un risque EUR/USD, "
   "que personne à Bamako ne pilote."),
  ("VAN — définition",
   "Somme des flux mensuels actualisés des mois 1 à 24, moins "
   "l'investissement du mois 0. Taux d'actualisation : 14 % par an, exigé "
   "par la direction, converti en taux mensuel équivalent."),
  ("TRI — définition",
   "Taux qui annule la VAN. Calculé sur les flux mensuels, mois 0 inclus, "
   "puis annualisé. Il n'a de sens que parce que les flux changent de "
   "signe une seule fois."),
  ("Taux limite — définition",
   "Taux de change de départ au-delà duquel la VAN devient négative, à "
   "dérive et volumes inchangés. Lu par Valeur cible, et recalculé par "
   "formule : la VAN vaut (valeur actuelle des encaissements CFA) / taux "
   "− (valeur actuelle des décaissements USD)."),
  ("Comment actualiser ce classeur",
   "Changer les valeurs bleues de HYPOTHESES, rien d'autre. Vérifier "
   "ensuite que les huit contrôles de QUALITE passent, et que "
   "REPONSES!B34 vaut toujours zéro. Une seule borne : H_MONTEE doit "
   "valoir au moins 2 — la courbe de montée en charge divise par "
   "(H_MONTEE − 1), et rendrait #DIV/0! à 1. Le cas n'a pas été "
   "corrigé : il a été écrit ici."),
  ("Les deux natures de flux, et pourquoi elles sont séparées",
   "CALCULS sépare volontairement la marge locale (encaissée en XOF, "
   "donc sensible au taux) des achats (décaissés en USD, donc "
   "insensibles). Ce n'est pas une redondance : c'est ce qui permet "
   "d'écrire la VAN sous la forme entrées/taux − sorties, et donc "
   "d'inverser la formule pour trouver le taux de bascule sans "
   "tâtonner."),
  ("Ce que ce modèle ne dit pas",
   "Il ne dit rien du besoin en fonds de roulement au-delà du stock "
   "initial, rien de la fiscalité malienne, rien du risque de sécurité. "
   "Trois trous connus, écrits ici pour qu'on ne les découvre pas en "
   "réunion."),
 ],
 "ANNEXE_IA": [
  ("Le prompt de cadrage",
   "Ce que vous avez demandé pour obtenir la liste des hypothèses à "
   "chiffrer, avant d'écrire la moindre formule."),
  ("Le prompt de structure",
   "Ce que vous avez demandé pour obtenir l'architecture du classeur."),
  ("L'erreur de l'IA",
   "Ce qu'elle a proposé et qui était faux ou incomplet, et comment vous "
   "l'avez attrapé. Au moins une erreur, décrite précisément."),
  ("Le V qui l'a attrapée",
   "V1 source · V2 calcul · V3 cohérence · V4 terrain · V5 lisibilité."),
  ("Ce que vous n'avez pas suivi",
   "Une suggestion de l'IA que vous avez écartée, et pourquoi."),
 ],
}

def feuille_texte(wb, nom, corrige):
    ws = wb.create_sheet(nom)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 104
    ws["A1"] = nom; ws["A1"].font = Ftitre
    for k, (cle, val) in enumerate(TEXTES[nom], start=3):
        c = ws.cell(k, 1, cle); c.font, c.border = Fgras, Bord
        d = ws.cell(k, 2, val if corrige else ""); d.border = Bord
        d.alignment = Alignment(wrap_text=True, vertical="top")
        if not corrige: d.fill = Rjaune
        ws.row_dimensions[k].height = 46
    return ws

# ── Les plages nommées ────────────────────────────────────────────────
PLAGES = (
 [(c, f"HYPOTHESES!$B${5+k}") for k, (c, _v, _f, _d) in enumerate(H)]
 + [("F_TAUX",        "CALCULS!$B$7:$B$30"),
    ("F_VOLUME",      "CALCULS!$C$7:$C$30"),
    ("F_CA_USD",      "CALCULS!$F$7:$F$30"),
    ("F_CHARGES_USD", "CALCULS!$G$7:$G$30"),
    ("F_MARGE_USD",   "CALCULS!$H$7:$H$30"),
    ("F_ACHATS",      "CALCULS!$I$7:$I$30"),
    ("F_FLUX",        "CALCULS!$J$6:$J$30"),
    ("F_FLUX1",       "CALCULS!$J$7:$J$30"),
    ("F_ACTU",        "CALCULS!$K$7:$K$30"),
    ("F_CUMUL",       "CALCULS!$M$6:$M$30"),
    ("S_FLUX_PESS",   "ANALYSE!$B$7:$B$30"),
    ("S_FLUX_OPT",    "ANALYSE!$D$7:$D$30"),
    ("Z_AUDIT",       "CALCULS!$A$6:$M$30"),
    ("R_ECART_TAUX",  "REPONSES!$B$35"),
    ("R_CONTROLE",    "REPONSES!$B$36")]
)
#  Les zones de construction libre : le correcteur ne note pas leurs
#  cellules une à une. Deux projections justes n'ont pas forcément les
#  mêmes colonnes ; ce qu'on note, ce sont les réponses qu'elles nourrissent.
ZONES_LIBRES = [
 ("LIBRE_CLEAN",     "CLEAN!$A$1:$H$40"),
 ("LIBRE_CALCULS",   "CALCULS!$A$1:$M$30"),
 ("LIBRE_ANALYSE",   "ANALYSE!$A$1:$U$60"),
 ("LIBRE_DASHBOARD", "DASHBOARD!$A$1:$H$40"),
]

def classeur_tp(corrige):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    feuille_texte(wb, "README", corrige)
    feuille_hypotheses(wb)
    feuille_raw(wb)
    feuille_clean(wb, True)          # donnée dans les deux : c'est du contexte
    feuille_calculs(wb, corrige)
    feuille_analyse(wb, corrige)
    feuille_dashboard(wb, corrige)
    feuille_qualite(wb, corrige)
    feuille_reponses(wb, corrige)
    feuille_texte(wb, "ANNEXE_IA", corrige)
    if corrige:
        for nom, cible in PLAGES + ZONES_LIBRES:
            wb.defined_names[nom] = DefinedName(nom, attr_text=cible)
    nom = "TP09_CORRIGE.xlsx" if corrige else "TP09_DEPART.xlsx"
    wb.save(RACINE / nom)
    print(f"  {nom}")

def classeur_jour():
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    feuille_texte(wb, "README", True)
    feuille_hypotheses(wb)
    feuille_raw(wb)
    feuille_clean(wb, True)
    wb.save(RACINE / "J09_BusinessPlan_7e_agence.xlsx")
    print("  J09_BusinessPlan_7e_agence.xlsx")

classeur_jour()
classeur_tp(False)
classeur_tp(True)
json.dump(R, open(RACINE / "REFERENCES_M09.json", "w"), indent=1, ensure_ascii=False)
print("  REFERENCES_M09.json")
