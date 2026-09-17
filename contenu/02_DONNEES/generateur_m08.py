#!/usr/bin/env python3
# ══════════════════════════════════════════════════════════════════════
#  Données du module 8 — le schéma en étoile
#
#  Une table de faits de 1,2 million de lignes et cinq dimensions.
#  La table de faits ne porte QUE des clés et des mesures : c'est ce
#  qui la rend compressible, et c'est tout le sujet du module.
#
#  Deux années civiles pleines — 2025 et 2026 — pour que la
#  « time intelligence » ait de quoi comparer.
# ══════════════════════════════════════════════════════════════════════
import datetime as dt, pathlib, sys, json, random, csv, math
from collections import defaultdict

random.seed(808)
RACINE = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "sortie-m08")
MODELE = RACINE / "J08_Modele"
MODELE.mkdir(parents=True, exist_ok=True)

DEBUT, FIN = dt.date(2025, 1, 1), dt.date(2026, 12, 31)
EPOQUE = dt.date(1899, 12, 30)          # le zéro des dates d'Excel
N_LIGNES = 1_200_000

# ══ 1 · d_Agence ══════════════════════════════════════════════════════
AGENCES = [
    (1, "Kinshasa",   "RD Congo",     "CDF", "Nadège Kalala",  0.27, 0.0055),
    (2, "Lubumbashi", "RD Congo",     "CDF", "Trésor Ilunga",  0.16, 0.0030),
    (3, "Abidjan",    "Côte d'Ivoire","XOF", "Serge Kouadio",  0.17, 0.0165),
    (4, "Dakar",      "Sénégal",      "XOF", "Aïcha Ndiaye",   0.13, 0.0075),
    (5, "Douala",     "Cameroun",     "XAF", "Thomas Mbarga",  0.16, 0.0025),
    (6, "Libreville", "Gabon",        "XAF", "Josée Nguema",   0.11, 0.0045),
]
DEVISE = {a[0]: a[3] for a in AGENCES}

# ══ 2 · d_Produit ═════════════════════════════════════════════════════
FAMILLES = {
    "Huiles":    (["Huile végétale", "Huile de palme", "Huile d'arachide"], 21.8),
    "Riz":       (["Riz parfumé", "Riz brisé", "Riz étuvé"],                44.8),
    "Farine":    (["Farine de blé", "Farine de maïs", "Farine de manioc"],  29.5),
    "Boissons":  (["Jus concentré", "Eau minérale", "Boisson gazeuse"],     14.1),
    "Savons":    (["Savon de ménage", "Détergent", "Savon de toilette"],    11.2),
    "Conserves": (["Tomate concentrée", "Sardines", "Haricots"],            18.8),
}
PART_CA = {"Huiles": 0.21, "Riz": 0.27, "Farine": 0.17,
           "Boissons": 0.12, "Savons": 0.10, "Conserves": 0.13}
CALIBRES = {"Huiles": ["1 L", "2 L", "5 L", "20 L"],
            "Riz": ["5 kg", "10 kg", "25 kg", "50 kg"],
            "Farine": ["1 kg", "5 kg", "25 kg", "50 kg"],
            "Boissons": ["33 cl", "50 cl", "1 L", "pack 6"],
            "Savons": ["lot 6", "lot 12", "lot 24", "lot 48"],
            "Conserves": ["70 g", "125 g", "400 g", "800 g"]}
#  Les taux de marge sont une donnée de la comptabilité, pas des ventes.
MARGE = {"Huiles": 0.172, "Riz": 0.118, "Farine": 0.143,
         "Boissons": 0.221, "Savons": 0.264, "Conserves": 0.195}

poids_lignes = {f: PART_CA[f] / FAMILLES[f][1] for f in FAMILLES}
t = sum(poids_lignes.values())
poids_lignes = {f: v / t for f, v in poids_lignes.items()}

produits = []
for i in range(1, 981):
    fam = random.choices(list(FAMILLES), weights=[poids_lignes[f] for f in FAMILLES])[0]
    base, pu = FAMILLES[fam]
    sous = random.choice(base)
    cal = random.choice(CALIBRES[fam])
    produits.append({
        "cle": i, "ref": f"REF-{i:04d}", "produit": f"{sous} {cal}",
        "famille": fam, "sous_famille": sous,
        "pu_usd": round(pu * 0.903 * random.uniform(0.55, 1.75), 2),
        #  Le taux de marge vit dans la DIMENSION, pas dans la table de
        #  faits. Le répéter sur 1,2 million de lignes coûterait neuf
        #  mégaoctets pour six valeurs distinctes.
        "taux_marge": MARGE[fam],
    })
PAR_FAMILLE = defaultdict(list)
for p in produits:
    PAR_FAMILLE[p["famille"]].append(p)

# ══ 3 · d_Client ══════════════════════════════════════════════════════
TETES = ["BOUTIQUE", "DEPOT", "SUPERETTE", "ALIMENTATION", "GROSSISTE",
         "MINI-MARCHE", "EPICERIE", "CANTINE", "RESTAURANT", "HOTEL",
         "BOULANGERIE", "SUPERMARCHE", "KIOSQUE", "MAGASIN"]
QUARTIERS = ["NGALULA", "KINTAMBO", "BANDAL", "MATONGE", "LEMBA", "GOMBE",
             "MASINA", "LIMETE", "SELEMBAO", "NGIRI", "VICTOIRE", "ABOBO",
             "PLATEAU", "SANDAGA", "BONABERI", "AKANDA", "TREICHVILLE",
             "YOPOUGON", "MEDINA", "BONAPRISO", "KATUBA", "RUASHI",
             "MAKELEKELE", "OWENDO", "PIKINE", "COCODY", "DEIDO", "LIKASI"]
SEGMENTS = ["Grossiste", "Détaillant", "Restauration", "Institution"]
POIDS_SEG = [0.18, 0.56, 0.18, 0.08]

clients = []
clients_par_agence = defaultdict(list)
cle = 0
for (ac, nom, pays, dev, resp, part, _c) in AGENCES:
    cible = round(3600 * part)
    vus = set()
    while len(vus) < cible:
        n = f"{random.choice(TETES)} {random.choice(QUARTIERS)} {len(vus) % 40 + 1:02d}"
        if n in vus:
            continue
        vus.add(n)
        cle += 1
        seg = random.choices(SEGMENTS, weights=POIDS_SEG)[0]
        clients.append({"cle": cle, "client": n, "segment": seg,
                        "agence_cle": ac, "ville": nom})
        clients_par_agence[ac].append(cle)
POIDS_CLIENT = {c["cle"]: {"Grossiste": 6.0, "Détaillant": 1.0,
                           "Restauration": 1.8, "Institution": 3.2}[c["segment"]]
                for c in clients}

#  Tous les clients du fichier n'achètent pas toutes les années : 8 %
#  arrivent en 2026, 6 % s'arrêtent fin 2025. Sans ça, DISTINCTCOUNT
#  rendrait 3 600 les deux années et la mesure n'apprendrait rien.
melange = [c["cle"] for c in clients]
random.shuffle(melange)
n = len(melange)
NOUVEAUX_2026 = set(melange[: int(n * 0.08)])
PARTIS_FIN_2025 = set(melange[int(n * 0.08): int(n * 0.14)])
for c in clients:
    c["arrivee"] = 2026 if c["cle"] in NOUVEAUX_2026 else 2025
    c["depart"] = 2025 if c["cle"] in PARTIS_FIN_2025 else 2026

# ══ 4 · d_Calendrier ══════════════════════════════════════════════════
FERIES = {(1, 1): "Nouvel An", (4, 1): "Martyrs de l'indépendance",
          (16, 1): "Laurent-Désiré Kabila", (17, 1): "Patrice Lumumba",
          (1, 5): "Fête du travail", (30, 6): "Indépendance",
          (1, 8): "Fête des parents", (25, 12): "Noël"}
MOIS_FR = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet",
           "août", "septembre", "octobre", "novembre", "décembre"]
MOIS_AB = ["janv.", "févr.", "mars", "avr.", "mai", "juin", "juil.",
           "août", "sept.", "oct.", "nov.", "déc."]
JOURS_FR = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]

#  Le calendrier couvre 2024 à 2026 : une table de dates doit couvrir
#  TOUT l'intervalle, du 1er janvier de la première année au 31 décembre
#  de la dernière. Un trou, et les cumuls annuels sont faux.
CAL_DEBUT, CAL_FIN = dt.date(2024, 1, 1), dt.date(2026, 12, 31)
calendrier = []
d = CAL_DEBUT
while d <= CAL_FIN:
    iso = d.isocalendar()
    ferie = FERIES.get((d.day, d.month))
    we = d.weekday() >= 5
    calendrier.append({
        "cle": (d - EPOQUE).days, "date": d, "annee": d.year,
        "trimestre": f"T{(d.month - 1)//3 + 1}", "mois_num": d.month,
        "mois_nom": MOIS_FR[d.month - 1], "mois_abrege": f"{MOIS_AB[d.month-1]} {d.year % 100:02d}",
        "debut_mois": dt.date(d.year, d.month, 1),
        "rang_mois": d.year * 12 + d.month,
        "semaine_iso": iso[1], "jour_semaine": d.weekday() + 1,
        "jour_nom": JOURS_FR[d.weekday()],
        "est_weekend": "VRAI" if we else "FAUX",
        "ferie": ferie or "",
        "est_ouvre": "FAUX" if (we or ferie) else "VRAI",
    })
    d += dt.timedelta(days=1)

# ══ 5 · d_Taux ════════════════════════════════════════════════════════
taux = {}
cdf, xof, xaf = 2890.0, 599.0, 599.0
m = dt.date(2024, 1, 1)
while m <= dt.date(2026, 12, 1):
    cdf *= 1 + random.uniform(0.002, 0.014)
    xof *= 1 + random.uniform(-0.004, 0.006)
    xaf *= 1 + random.uniform(-0.004, 0.006)
    taux[(m, "CDF")] = round(cdf, 2)
    taux[(m, "XOF")] = round(xof, 2)
    taux[(m, "XAF")] = round(xaf, 2)
    m = dt.date(m.year + (m.month == 12), m.month % 12 + 1, 1)

print(f"  dimensions : {len(AGENCES)} agences · {len(produits)} produits · "
      f"{len(clients)} clients · {len(calendrier)} jours · {len(taux)} taux")

# ══ 6 · f_Ventes — 1,2 million de lignes ══════════════════════════════
#  La granularité est LE choix du module : ici, une ligne = une ligne de
#  ticket. Pas une commande, pas un mois. La table de faits ne porte que
#  des clés et des mesures — aucun libellé. C'est ce qui la compresse.
SAISON = {1: 0.92, 2: 0.90, 3: 1.02, 4: 1.06, 5: 1.01, 6: 0.97,
          7: 0.99, 8: 1.07, 9: 1.00, 10: 1.03, 11: 1.08, 12: 1.21}
JOURS = [c for c in calendrier if DEBUT <= c["date"] <= FIN]
POIDS_AG = {a[0]: a[5] for a in AGENCES}
CROISS = {a[0]: a[6] for a in AGENCES}

#  Le poids de chaque jour : saison × croissance × jour ouvré.
poids_jour = {}
for c in JOURS:
    rang = (c["date"].year - 2025) * 12 + c["date"].month - 1
    base = SAISON[c["date"].month] * (0.55 if c["est_weekend"] == "VRAI" else 1.0)
    poids_jour[c["cle"]] = (base, rang)

lignes_par_agence = {a[0]: int(N_LIGNES * POIDS_AG[a[0]]) for a in AGENCES}
reste = N_LIGNES - sum(lignes_par_agence.values())
lignes_par_agence[1] += reste

JOUR_ANNEE_PRE = {c["cle"]: c["annee"] for c in calendrier}
CLI_ARRIVEE = {c["cle"]: c["arrivee"] for c in clients}
CLI_DEPART = {c["cle"]: c["depart"] for c in clients}

f_ventes = []
n_commande = 0
for ac, n_cible in lignes_par_agence.items():
    cles_jours = [c["cle"] for c in JOURS]
    poids = [poids_jour[k][0] * (1 + CROISS[ac] * poids_jour[k][1]) for k in cles_jours]
    cls = clients_par_agence[ac]
    poids_cl = [POIDS_CLIENT[c] for c in cls]
    ecrit = 0
    while ecrit < n_cible:
        jour = random.choices(cles_jours, weights=poids)[0]
        an_jour = JOUR_ANNEE_PRE[jour]
        client = random.choices(cls, weights=poids_cl)[0]
        if not (CLI_ARRIVEE[client] <= an_jour <= CLI_DEPART[client]):
            continue                      # ce client n'était pas actif cette année-là
        n_commande += 1
        #  Un ticket porte de 1 à 4 lignes : c'est ce qui fait que le
        #  nombre de commandes distinctes n'est pas le nombre de lignes.
        taille = min(n_cible - ecrit, random.choices([1, 2, 3, 4],
                                                    weights=[42, 30, 18, 10])[0])
        for _ in range(taille):
            fam = random.choices(list(FAMILLES),
                                 weights=[poids_lignes[f] for f in FAMILLES])[0]
            art = random.choice(PAR_FAMILLE[fam])
            qte = max(1, min(9, int(round(random.lognormvariate(0.55, 0.7)))))
            f_ventes.append((jour, ac, client, art["cle"], n_commande, qte,
                             art["pu_usd"], fam))
            ecrit += 1

print(f"  f_Ventes : {len(f_ventes):,} lignes · {n_commande:,} tickets"
      .replace(",", " "))

# ── Conversion en devise locale, et coût ─────────────────────────────
JOUR_MOIS = {c["cle"]: c["debut_mois"] for c in calendrier}
JOUR_ANNEE = {c["cle"]: c["annee"] for c in calendrier}
lignes_csv = []
ref = {"ca_usd": defaultdict(float), "marge_usd": defaultdict(float),
       "qte": defaultdict(int), "lignes": defaultdict(int)}
ref_ag = defaultdict(float)
ref_fam = defaultdict(float)
ref_mois = defaultdict(float)
clients_annee = defaultdict(set)
tickets_annee = defaultdict(set)

for (jour, ac, client, pcle, ncmd, qte, pu_usd, fam) in f_ventes:
    mois = JOUR_MOIS[jour]
    an = JOUR_ANNEE[jour]
    tx = taux[(mois, DEVISE[ac])]
    #  Ni le franc congolais ni le franc CFA n'ont de subdivision en
    #  circulation : les montants sont entiers, et ça retire une colonne
    #  de décimales sur 1,2 million de lignes.
    montant_local = int(round(pu_usd * qte * tx))
    usd = montant_local / tx
    lignes_csv.append((jour, ac, client, pcle, ncmd, qte, montant_local))
    ref["ca_usd"][an] += usd
    ref["marge_usd"][an] += montant_local * MARGE[fam] / tx
    ref["qte"][an] += qte
    ref["lignes"][an] += 1
    ref_ag[(an, ac)] += usd
    ref_fam[(an, fam)] += usd
    ref_mois[(an, mois.month)] += usd
    clients_annee[an].add(client)
    tickets_annee[an].add(ncmd)

for an in (2025, 2026):
    print(f"  {an} : CA {ref['ca_usd'][an]:>14,.2f} USD · "
          f"{len(clients_annee[an])} clients · {len(tickets_annee[an]):,} tickets"
          .replace(",", " "))

# ══ 7 · L'écriture des six tables ═════════════════════════════════════
#  Tout en CSV UTF-8, séparé par des points-virgules, décimale par le
#  point. Power Query lit ça sans discuter, et un CSV de 1,2 million de
#  lignes se compresse là où un .xlsx pèserait soixante mégaoctets.
def ecrire(nom, entetes, lignes):
    chemin = MODELE / nom
    with open(chemin, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(entetes)
        w.writerows(lignes)
    return chemin.stat().st_size

t1 = ecrire("f_Ventes.csv",
            ["Date_Cle", "Agence_Cle", "Client_Cle", "Produit_Cle",
             "Ticket_Cle", "Quantite", "Montant_Local"],
            lignes_csv)
t2 = ecrire("d_Agence.csv", ["Agence_Cle", "Agence", "Pays", "Devise", "Responsable"],
            [(a[0], a[1], a[2], a[3], a[4]) for a in AGENCES])
t3 = ecrire("d_Produit.csv",
            ["Produit_Cle", "Reference", "Produit", "Famille", "Sous_Famille",
             "PU_USD", "Taux_Marge"],
            [(p["cle"], p["ref"], p["produit"], p["famille"], p["sous_famille"],
              p["pu_usd"], p["taux_marge"]) for p in produits])
t4 = ecrire("d_Client.csv", ["Client_Cle", "Client", "Segment", "Agence_Cle", "Ville"],
            [(c["cle"], c["client"], c["segment"], c["agence_cle"], c["ville"])
             for c in clients])
t5 = ecrire("d_Calendrier.csv",
            ["Date_Cle", "Date", "Annee", "Trimestre", "Mois_Num", "Mois_Nom",
             "Mois_Abrege", "Debut_Mois", "Rang_Mois", "Semaine_ISO",
             "Jour_Semaine", "Jour_Nom", "Est_WeekEnd", "Ferie", "Est_Ouvre"],
            [(c["cle"], c["date"].isoformat(), c["annee"], c["trimestre"],
              c["mois_num"], c["mois_nom"], c["mois_abrege"],
              c["debut_mois"].isoformat(), c["rang_mois"], c["semaine_iso"],
              c["jour_semaine"], c["jour_nom"], c["est_weekend"], c["ferie"],
              c["est_ouvre"]) for c in calendrier])
#  d_Taux porte UNE ligne par mois, et une colonne par devise. C'est la
#  seule forme qui lui donne une clé unique — donc une relation. Avec
#  trois lignes par mois, Power Pivot refuse la relation, et la
#  conversion se fait alors ligne à ligne, ce qui est lent et faux.
mois_taux = sorted({m for (m, _d) in taux})
t6 = ecrire("d_Taux.csv", ["Debut_Mois", "CDF", "XOF", "XAF"],
            [(m.isoformat(), taux[(m, "CDF")], taux[(m, "XOF")], taux[(m, "XAF")])
             for m in mois_taux])

import zipfile
zchemin = RACINE / "J08_Modele.zip"
with zipfile.ZipFile(zchemin, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
    for p in sorted(MODELE.iterdir()):
        z.write(p, f"J08_Modele/{p.name}")
print(f"  f_Ventes.csv {t1/1e6:.1f} Mo · dimensions {(t2+t3+t4+t5+t6)/1e3:.0f} Ko")
print(f"  J08_Modele.zip {zchemin.stat().st_size/1e6:.1f} Mo")

# ══ 8 · Les valeurs de référence ══════════════════════════════════════
#  Power Pivot n'existe pas sur macOS : le DAX de ce module est écrit
#  mais ne peut pas être exécuté ici. Ces valeurs sont donc calculées en
#  Python, sur les mêmes données que celles qui sont livrées, et ce sont
#  elles qui font foi.
AN = 2026
R = {}
R["lignes_faits"] = len(lignes_csv)
R["tickets"] = n_commande
R["tables"] = 6
R["relations"] = 5
R["ca_2026"] = round(ref["ca_usd"][2026], 2)
R["ca_2025"] = round(ref["ca_usd"][2025], 2)
R["evolution"] = round(ref["ca_usd"][2026] / ref["ca_usd"][2025] - 1, 6)
R["marge_pct_2026"] = round(ref["marge_usd"][2026] / ref["ca_usd"][2026], 6)
R["clients_2026"] = len(clients_annee[2026])
R["clients_2025"] = len(clients_annee[2025])
R["tickets_2026"] = len(tickets_annee[2026])
R["panier_2026"] = round(ref["ca_usd"][2026] / len(tickets_annee[2026]), 2)
R["ca_kinshasa_2026"] = round(ref_ag[(2026, 1)], 2)
R["ca_par_agence_2026"] = {a[1]: round(ref_ag[(2026, a[0])], 2) for a in AGENCES}
R["ca_par_famille_2026"] = {f: round(ref_fam[(2026, f)], 2)
                            for f in sorted(FAMILLES, key=lambda x: -ref_fam[(2026, x)])}
R["cumul_juin_2026"] = round(sum(ref_mois[(2026, m)] for m in range(1, 7)), 2)
R["ca_par_mois_2026"] = {m: round(ref_mois[(2026, m)], 2) for m in range(1, 13)}
R["taille_csv_mo"] = round(t1 / 1e6, 1)
R["taille_zip_mo"] = round(zchemin.stat().st_size / 1e6, 1)

json.dump(R, open(RACINE / "REFERENCES_M08.json", "w"), indent=1, ensure_ascii=False)
for k, v in R.items():
    if not isinstance(v, dict):
        print(f"  {k:<20} {v}")

# ══ 9 · Le travail pratique ═══════════════════════════════════════════
import openpyxl, zipfile as _zip, shutil, re as _re
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

Ftitre  = Font(name="Calibri", size=14, bold=True)
Fentete = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
Faide   = Font(name="Calibri", size=10, italic=True, color="FF6B7280")
Fgras   = Font(name="Calibri", size=11, bold=True)
Fmono   = Font(name="Consolas", size=10)
Rentete = PatternFill("solid", fgColor="FF243044")
Rjaune  = PatternFill("solid", fgColor="FFFFF6D8")
Rvert   = PatternFill("solid", fgColor="FFE8F6EE")
Bord    = Border(bottom=Side(style="thin", color="FFD6DBE1"))
FMT_USD, FMT_PCT, FMT_INT = '#,##0.00\\ "USD"', "0.00%", "#,##0"

def entete(ws, cols, ligne=1, depart=1):
    for i, c in enumerate(cols, start=depart):
        cel = ws.cell(ligne, i, c)
        cel.font, cel.fill = Fentete, Rentete
        cel.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[ligne].height = 22

def colonnes(ws, specs, depart=1):
    for i, (w, fmt) in enumerate(specs, start=depart):
        d = ws.column_dimensions[get_column_letter(i)]
        d.width = w
        if fmt: d.number_format = fmt

def tableau(ws, nom, ref):
    t = Table(displayName=nom, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)

MOD = "ThisWorkbookDataModel"
def cube(mesure, *membres):
    args = ",".join([f'"{MOD}"', f'"[Measures].[{mesure}]"'] + [f'"{m}"' for m in membres])
    return f"CUBEVALUE({args})"

AN26 = "[d_Calendrier].[Annee].&[2026]"

#  Les douze réponses. La valeur attendue accompagne la formule : elle
#  est calculée en Python, et c'est elle qu'on injectera dans le cache
#  du corrigé — Power Pivot n'existant pas sur macOS, Excel ne peut pas
#  évaluer un CUBEVALEUR ici.
REPONSES = [
 ("Chiffre d'affaires 2026, en USD",
  f"={cube('CA USD', AN26)}", FMT_USD, R["ca_2026"]),
 ("Chiffre d'affaires 2025, par la mesure N-1",
  f"={cube('CA USD N-1', AN26)}", FMT_USD, R["ca_2025"]),
 ("Évolution 2026 / 2025", "=B10/B11-1", FMT_PCT, R["evolution"]),
 ("Taux de marge 2026", f"={cube('Marge %', AN26)}", FMT_PCT, R["marge_pct_2026"]),
 ("Clients actifs en 2026", f"={cube('Clients actifs', AN26)}", FMT_INT, R["clients_2026"]),
 ("Chiffre d'affaires cumulé à fin juin 2026",
  f"={cube('CA USD Cumul Annuel', AN26, '[d_Calendrier].[Mois_Num].&[6]')}",
  FMT_USD, R["cumul_juin_2026"]),
 ("Chiffre d'affaires de Kinshasa en 2026",
  f"={cube('CA USD', AN26, '[d_Agence].[Agence].&[Kinshasa]')}",
  FMT_USD, R["ca_kinshasa_2026"]),
 ("Panier moyen par ticket, en USD",
  f"=B10/{cube('Nb Tickets', AN26)}", FMT_USD, R["panier_2026"]),
 ("Nombre de tables du modèle", "=COUNTA(t_Tables[Table])", FMT_INT, R["tables"]),
 ("Nombre de relations du modèle", "=COUNTA(t_Relations[Table_source])",
  FMT_INT, R["relations"]),
 ("R_ECART_MESURE — TCD moins CUBEVALEUR  ·  doit valoir 0", "=B5-B10", FMT_USD, 0.0),
 ("R_CONTROLE — total moins la somme des six agences  ·  doit valoir 0",
  "=B10-(" + "+".join(cube("CA USD", AN26, f"[d_Agence].[Agence].&[{a[1]}]")
                      for a in AGENCES) + ")", FMT_USD, 0.0),
]

TABLES_MODELE = [
 ("f_Ventes", "Faits", 1_200_000, "aucune — c'est la table de faits",
  "une ligne = une ligne de ticket"),
 ("d_Calendrier", "Dimension", 1_096, "Date_Cle",
  "obligatoire, et à marquer comme table de dates"),
 ("d_Agence", "Dimension", 6, "Agence_Cle", "porte la devise de l'agence"),
 ("d_Client", "Dimension", 3_600, "Client_Cle", "porte le segment"),
 ("d_Produit", "Dimension", 980, "Produit_Cle", "porte le taux de marge"),
 ("d_Taux", "Dimension", 36, "Debut_Mois",
  "une ligne par mois, une colonne par devise — c'est ce qui lui donne une clé unique"),
]
RELATIONS_MODELE = [
 ("f_Ventes", "Date_Cle", "d_Calendrier", "Date_Cle", "plusieurs-à-un"),
 ("f_Ventes", "Agence_Cle", "d_Agence", "Agence_Cle", "plusieurs-à-un"),
 ("f_Ventes", "Client_Cle", "d_Client", "Client_Cle", "plusieurs-à-un"),
 ("f_Ventes", "Produit_Cle", "d_Produit", "Produit_Cle", "plusieurs-à-un"),
 ("d_Calendrier", "Debut_Mois", "d_Taux", "Debut_Mois", "plusieurs-à-un (flocon)"),
]
MESURES = [
 ("CA Local", "SUM ( f_Ventes[Montant_Local] )"),
 ("CA USD",
  "SUMX ( f_Ventes ; DIVIDE ( f_Ventes[Montant_Local] ; "
  "SWITCH ( RELATED ( d_Agence[Devise] ) ; "
  '"CDF" ; RELATED ( d_Taux[CDF] ) ; '
  '"XOF" ; RELATED ( d_Taux[XOF] ) ; '
  '"XAF" ; RELATED ( d_Taux[XAF] ) ) ) )'),
 ("Marge USD",
  "SUMX ( f_Ventes ; DIVIDE ( f_Ventes[Montant_Local] * "
  "RELATED ( d_Produit[Taux_Marge] ) ; "
  "SWITCH ( RELATED ( d_Agence[Devise] ) ; "
  '"CDF" ; RELATED ( d_Taux[CDF] ) ; '
  '"XOF" ; RELATED ( d_Taux[XOF] ) ; '
  '"XAF" ; RELATED ( d_Taux[XAF] ) ) ) )'),
 ("Marge %", "DIVIDE ( [Marge USD] ; [CA USD] ; 0 )"),
 ("Clients actifs", "DISTINCTCOUNT ( f_Ventes[Client_Cle] )"),
 ("Nb Tickets", "DISTINCTCOUNT ( f_Ventes[Ticket_Cle] )"),
 ("CA USD N-1", "CALCULATE ( [CA USD] ; SAMEPERIODLASTYEAR ( d_Calendrier[Date] ) )"),
 ("CA USD Cumul Annuel", "TOTALYTD ( [CA USD] ; d_Calendrier[Date] )"),
]

def feuille_reponses(wb, corrige):
    ws = wb.create_sheet("REPONSES")
    ws.column_dimensions["A"].width = 64
    ws.column_dimensions["B"].width = 24
    ws["A1"] = "TP 8 — Le modèle qui remplace 42 RECHERCHEX"; ws["A1"].font = Ftitre
    ws["A2"] = ("Année analysée : 2026. Les cellules jaunes se remplissent une fois "
                "le modèle construit et les huit mesures écrites.")
    ws["A2"].font = Faide
    ws["A4"] = "①  CE QUE VOUS LISEZ DANS LE TABLEAU CROISÉ"; ws["A4"].font = Fgras
    ws["A5"] = "Chiffre d'affaires 2026 en USD, lu dans le TCD branché sur le modèle"
    ws["A5"].border = Bord
    c = ws.cell(5, 2); c.fill, c.border, c.number_format = Rjaune, Bord, FMT_USD
    if corrige:
        c.value = R["ca_2026"]
    ws["A8"] = "②  CE QUE CUBEVALEUR ET LES FORMULES VÉRIFIENT"; ws["A8"].font = Fgras
    for i, (lib, f, fmt, _v) in enumerate(REPONSES, start=10):
        ws.cell(i, 1, lib).border = Bord
        c = ws.cell(i, 2)
        c.fill, c.border, c.number_format = Rjaune, Bord, fmt
        if corrige:
            c.value = f
    for i in (20, 21):
        ws.cell(i, 1).font = Fgras
        ws.cell(i, 2).fill = Rvert
    return ws

def feuille_modele(wb, corrige):
    ws = wb.create_sheet("MODELE")
    ws["A1"] = "Le modèle, documenté"; ws["A1"].font = Ftitre
    ws["A2"] = ("Un modèle qu'on ne peut pas décrire en une page est un modèle "
                "que personne ne reprendra.")
    ws["A2"].font = Faide

    ws["A4"] = "LES TABLES"; ws["A4"].font = Fgras
    entete(ws, ["Table", "Rôle", "Lignes", "Clé", "Remarque"], ligne=5)
    for i, t in enumerate(TABLES_MODELE if corrige else [("", "", None, "", "")] * 6,
                          start=6):
        for j, v in enumerate(t, start=1):
            cel = ws.cell(i, j, v if v != "" else None)
            cel.border = Bord
            if not corrige: cel.fill = Rjaune
    tableau(ws, "t_Tables", "A5:E11")

    ws["A14"] = "LES RELATIONS"; ws["A14"].font = Fgras
    entete(ws, ["Table_source", "Colonne_source", "Table_cible", "Colonne_cible",
                "Cardinalite"], ligne=15)
    for i, t in enumerate(RELATIONS_MODELE if corrige else [("", "", "", "", "")] * 5,
                          start=16):
        for j, v in enumerate(t, start=1):
            cel = ws.cell(i, j, v if v != "" else None)
            cel.border = Bord
            if not corrige: cel.fill = Rjaune
    tableau(ws, "t_Relations", "A15:E20")

    ws["A23"] = "LES HUIT MESURES"; ws["A23"].font = Fgras
    entete(ws, ["Mesure", "Définition DAX"], ligne=24)
    for i, (nom, dax) in enumerate(MESURES if corrige else [("", "")] * 8, start=25):
        a = ws.cell(i, 1, nom or None); a.border = Bord
        b = ws.cell(i, 2, dax or None); b.border, b.font = Bord, Fmono
        b.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 46
        if not corrige:
            a.fill = b.fill = Rjaune
    tableau(ws, "t_Mesures", "A24:B32")
    colonnes(ws, [(16, None), (18, None), (12, FMT_INT), (16, None), (58, None)])
    ws.column_dimensions["B"].width = 96 if corrige else 18
    return ws

TEXTES = {
 "README": [
   ("Le modèle", "Schéma en étoile : une table de faits, cinq dimensions, "
                 "cinq relations plusieurs-à-un."),
   ("Granularité", "Une ligne de f_Ventes = une ligne de ticket. Pas une commande, "
                   "pas un mois. Tout le reste en découle."),
   ("La table de calendrier", "d_Calendrier couvre 2024 à 2026 en entier — du 1er "
                              "janvier de la première année au 31 décembre de la "
                              "dernière. Elle est MARQUÉE comme table de dates. "
                              "Sans ce clic, SAMEPERIODLASTYEAR rend vide."),
   ("d_Taux", "Une ligne par mois, une colonne par devise. C'est la seule forme qui "
              "lui donne une clé unique, donc une relation. Elle se relie à "
              "d_Calendrier par Debut_Mois : c'est un flocon, assumé."),
   ("Le taux de marge", "Il vit dans d_Produit, pas dans la table de faits. Six "
                        "valeurs distinctes répétées sur 1,2 million de lignes "
                        "coûteraient neuf mégaoctets pour rien."),
   ("Chargement", "Les six tables sont chargées en CONNEXION SEULE + modèle de "
                  "données. Aucune n'est écrite dans une feuille : f_Ventes ne "
                  "tiendrait pas, la limite d'une feuille est 1 048 576 lignes."),
   ("Poids", "39,1 Mo de CSV en entrée. Le classeur final pèse quelques mégaoctets : "
             "le modèle compresse par colonne et par valeur distincte."),
   ("Conversion", "SUMX + RELATED + SWITCH. La conversion se fait ligne à ligne au "
                  "taux du mois de la ligne — pas à un taux moyen."),
   ("Contrôles", "R_ECART_MESURE et R_CONTROLE doivent valoir 0 avant tout envoi."),
   ("Sur Mac", "Power Pivot n'existe ni sur macOS ni sur Excel pour le web. "
               "Le plan B est décrit dans l'énoncé."),
 ],
 "ANNEXE_IA": [
   ("Prompt 1", "Voici le schéma de mon modèle — les tables, leurs colonnes, leurs "
                "relations. Pas les données. Explique-moi ce que filtre quoi."),
   ("Prompt 2", "Écris-moi la mesure [nom] en DAX pour Power Pivot dans Excel 2024. "
                "Puis explique-la ligne par ligne, et donne-moi trois cas de test."),
   ("Prompt 3", "Ma mesure SAMEPERIODLASTYEAR rend vide. Voici mon modèle. "
                "Donne-moi les trois causes les plus probables, dans l'ordre, et "
                "comment vérifier chacune en dix secondes."),
   ("Erreur détectée", "L'IA a proposé une mesure utilisant une fonction DAX "
                       "récente, disponible dans Power BI mais absente de Power "
                       "Pivot : le moteur d'Excel ne suit pas le même rythme de "
                       "mise à jour. La formule est refusée à la saisie."),
   ("Quel V l'a attrapée", "V1 — Version. La question à poser dans le prompt : "
                           "« pour Power Pivot dans Excel 2024, pas pour Power BI »."),
 ],
}
QUALITE = [
 ("Q_TABLES — tables manquantes dans la documentation  ·  doit valoir 0",
  "=6-COUNTA(t_Tables[Table])", FMT_INT, 0,
  "six tables : une de faits, cinq dimensions"),
 ("Q_MESURES — mesures écrites  ·  doit valoir 8", "=COUNTA(t_Mesures[Mesure])",
  FMT_INT, 8, "huit mesures, nommées en français métier"),
 ("Q_RELATIONS — relations manquantes  ·  doit valoir 0",
  "=5-COUNTA(t_Relations[Table_source])", FMT_INT, 0,
  "cinq relations : quatre depuis f_Ventes, une en flocon vers d_Taux"),
]

def feuille_qualite(wb, corrige):
    ws = wb.create_sheet("QUALITE")
    colonnes(ws, [(54, None), (16, None), (62, None)])
    ws["A1"] = "Contrôles de qualité"; ws["A1"].font = Ftitre
    entete(ws, ["Contrôle", "Valeur", "Ce que ça veut dire"], ligne=3)
    for i, (lib, f, fmt, _v, quoi) in enumerate(QUALITE, start=4):
        ws.cell(i, 1, lib).border = Bord
        c = ws.cell(i, 2); c.fill, c.border, c.number_format = Rjaune, Bord, fmt
        if corrige: c.value = f
        ws.cell(i, 3, quoi).font = Faide
    return ws

def feuille_texte(wb, nom, corrige):
    ws = wb.create_sheet(nom)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100
    ws["A1"] = nom; ws["A1"].font = Ftitre
    for i, (cle, val) in enumerate(TEXTES[nom], start=3):
        c = ws.cell(i, 1, cle); c.font, c.border = Fgras, Bord
        d = ws.cell(i, 2, val if corrige else ""); d.border = Bord
        d.alignment = Alignment(wrap_text=True, vertical="top")
        if not corrige: d.fill = Rjaune
        ws.row_dimensions[i].height = 34
    return ws

def injecter_valeurs(chemin, valeurs):
    """openpyxl écrit une formule OU une valeur, jamais les deux — et
    Excel ne peut pas évaluer un CUBEVALEUR sans modèle, ce que macOS
    ne sait pas faire. On injecte donc le cache à la main, depuis les
    références calculées en Python."""
    chemin = pathlib.Path(chemin)
    tmp = chemin.with_suffix(".tmp.xlsx")
    poses = 0
    with _zip.ZipFile(chemin) as zin, _zip.ZipFile(tmp, "w", _zip.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            m = _re.match(r"xl/worksheets/sheet(\d+)\.xml$", it.filename)
            if m:
                txt = data.decode("utf-8")
                for ref, val in valeurs.get(int(m.group(1)), {}).items():
                    motif = _re.compile(
                        r'(<c r="' + ref + r'"[^>]*>)(<f>.*?</f>)<v>[^<]*</v>(</c>)')
                    txt, k = motif.subn(
                        lambda g: g.group(1) + g.group(2) + f"<v>{val!r}</v>" + g.group(3),
                        txt)
                    poses += k
                data = txt.encode("utf-8")
            zout.writestr(it, data)
    shutil.move(tmp, chemin)
    return poses

def classeur_tp(corrige):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    feuille_reponses(wb, corrige)
    feuille_modele(wb, corrige)
    feuille_texte(wb, "README", corrige)
    feuille_qualite(wb, corrige)
    feuille_texte(wb, "ANNEXE_IA", corrige)
    if corrige:
        for nom, cible in (("R_ECART_MESURE", "REPONSES!$B$20"),
                           ("R_CONTROLE", "REPONSES!$B$21"),
                           ("Q_MESURES", "QUALITE!$B$5")):
            wb.defined_names[nom] = DefinedName(nom, attr_text=cible)
    nom = "TP08_CORRIGE.xlsx" if corrige else "TP08_DEPART.xlsx"
    wb.save(RACINE / nom)
    if corrige:
        vals = {1: {f"B{10+i}": v for i, (_l, _f, _fm, v) in enumerate(REPONSES)},
                4: {f"B{4+i}": v for i, (_l, _f, _fm, v, _q) in enumerate(QUALITE)}}
        n = injecter_valeurs(RACINE / nom, vals)
        print(f"  {nom}  {n} valeurs injectées dans le cache")
    else:
        print(f"  {nom}")

classeur_tp(False)
classeur_tp(True)
