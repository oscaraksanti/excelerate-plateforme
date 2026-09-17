#!/usr/bin/env python3
# ══════════════════════════════════════════════════════════════════════
#  Données du module 10 — automatiser & livrer
#
#  200 classeurs clients au même format, à traiter en lot. « Au même
#  format » est le mot important : c'est ce qui rend le traitement
#  automatisable. Les quinze qui n'y sont pas tout à fait sont
#  délibérés — c'est sur eux que se joue la différence entre un script
#  qui tourne et un script qui tient.
#
#  Période : août 2026. Les taux de change sont ceux du mois, figés.
# ══════════════════════════════════════════════════════════════════════
import datetime as dt, pathlib, sys, json, random, zipfile, shutil

random.seed(1010)
RACINE = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "sortie-m10")
CLIENTS_DIR = RACINE / "J10_Clients"
if CLIENTS_DIR.exists():
    shutil.rmtree(CLIENTS_DIR)
CLIENTS_DIR.mkdir(parents=True, exist_ok=True)

DEBUT, FIN = dt.date(2026, 8, 1), dt.date(2026, 8, 31)

# ══ 1 · Le référentiel ════════════════════════════════════════════════
AGENCES = [
    ("Kinshasa",   "RD Congo",      "CDF", 0.27),
    ("Lubumbashi", "RD Congo",      "CDF", 0.16),
    ("Abidjan",    "Côte d'Ivoire", "XOF", 0.17),
    ("Dakar",      "Sénégal",       "XOF", 0.13),
    ("Douala",     "Cameroun",      "XAF", 0.16),
    ("Libreville", "Gabon",         "XAF", 0.11),
]
#  XOF et XAF sont tous deux arrimés à l'euro à la même parité
#  (655,957 pour 1 EUR) : leur taux face au dollar est donc identique,
#  au centime près. Ce n'est pas une coquille, c'est la définition.
TAUX = {"CDF": 3145.80, "XOF": 648.20, "XAF": 648.20}

FAMILLES = {
    "Huiles":    [("Huile végétale", "5 L", 6.40), ("Huile de palme", "20 L", 22.10),
                  ("Huile d'arachide", "2 L", 3.95)],
    "Riz":       [("Riz parfumé", "25 kg", 24.80), ("Riz brisé", "50 kg", 38.50),
                  ("Riz étuvé", "10 kg", 11.20)],
    "Farine":    [("Farine de blé", "50 kg", 29.40), ("Farine de maïs", "25 kg", 13.60),
                  ("Farine de manioc", "5 kg", 3.10)],
    "Boissons":  [("Jus concentré", "pack 6", 8.70), ("Eau minérale", "pack 6", 2.30),
                  ("Boisson gazeuse", "pack 6", 4.15)],
    "Savons":    [("Savon de ménage", "lot 24", 9.80), ("Détergent", "lot 12", 14.60),
                  ("Savon de toilette", "lot 48", 18.20)],
    "Conserves": [("Tomate concentrée", "400 g", 1.05), ("Sardines", "125 g", 0.88),
                  ("Haricots", "800 g", 1.72)],
}
PART_FAMILLE = {"Huiles": 0.21, "Riz": 0.27, "Farine": 0.17,
                "Boissons": 0.12, "Savons": 0.10, "Conserves": 0.13}

PREFIXES = ["Établissements", "Société", "Comptoir", "Maison", "Entreprise",
            "Ets", "Groupe", "Alimentation", "Distribution", "Négoce"]
NOMS = ["Kalala", "Ilunga", "Kouadio", "Ndiaye", "Mbarga", "Nguema", "Mutombo",
        "Traoré", "Diallo", "Bakayoko", "Essomba", "Ondo", "Kabila", "Sylla",
        "Camara", "Fall", "Sarr", "Owona", "Mba", "Tshibangu", "Mukendi",
        "Koffi", "Yao", "Bamba", "Cissé", "Gueye", "Diop", "Ndong", "Ntoutoume",
        "Biya", "Atangana", "Kamga", "Fotso", "Nzé", "Moussavou", "Lubaki",
        "Mwanza", "Kasongo", "Bemba", "Tchicaya"]
SUFFIXES = ["SARL", "SA", "& Fils", "SARLU", "Sarl", "", "", "SA"]

def nom_client(k):
    random.seed(1010 + k * 7)
    p, n, s = random.choice(PREFIXES), random.choice(NOMS), random.choice(SUFFIXES)
    return f"{p} {n} {s}".strip().replace("  ", " ")

# ══ 2 · Les 200 clients, et les quinze anomalies délibérées ═══════════
#  Quatre vides · six avec une ligne TOTAL en bas · trois avec une
#  colonne en plus · deux dont la feuille n'a pas la même casse.
random.seed(1010)
clients = []
for k in range(1, 201):
    ag = random.choices(AGENCES, weights=[a[3] for a in AGENCES])[0]
    clients.append({
        "k": k, "code": f"C{k:04d}", "nom": nom_client(k),
        "agence": ag[0], "pays": ag[1], "devise": ag[2],
        "poids": random.lognormvariate(0, 0.62),
    })

VIDES        = {17, 58, 129, 184}
AVEC_TOTAL   = {6, 41, 77, 112, 150, 193}
COLONNE_PLUS = {23, 96, 168}
CASSE_AUTRE  = {34, 141}

for c in clients:
    k = c["k"]
    c["anomalie"] = ("vide" if k in VIDES else
                     "ligne TOTAL en bas" if k in AVEC_TOTAL else
                     "colonne Remise_Pct en plus" if k in COLONNE_PLUS else
                     "feuille COMMANDES en majuscules" if k in CASSE_AUTRE else "")

# ══ 3 · Les commandes ═════════════════════════════════════════════════
JOURS = [DEBUT + dt.timedelta(days=d) for d in range((FIN - DEBUT).days + 1)]
JOURS = [j for j in JOURS if j.weekday() < 6]      # fermé le dimanche

lignes = []          # la consolidation attendue
for c in clients:
    if c["anomalie"] == "vide":
        c["lignes"], c["total_local"] = 0, 0.0
        c["total_usd"] = c["total_usd_direct"] = 0.0
        c["commandes"] = []
        continue
    n = max(4, int(random.gauss(30, 11) * c["poids"] ** 0.5))
    n = min(n, 78)
    cmds = []
    for _ in range(n):
        fam = random.choices(list(FAMILLES), weights=[PART_FAMILLE[f] for f in FAMILLES])[0]
        prod, cal, pu_usd = random.choice(FAMILLES[fam])
        jour = random.choice(JOURS)
        qte = max(1, int(random.lognormvariate(2.4, 0.85) * c["poids"]))
        tx = TAUX[c["devise"]]
        #  Le prix local est arrondi comme il l'est sur une facture :
        #  au franc près, jamais au centime — ni le CDF ni le CFA n'ont
        #  de subdivision en circulation.
        pu_local = round(pu_usd * tx * random.uniform(0.97, 1.05))
        montant_local = pu_local * qte
        cmds.append({
            "date": jour, "reference": f"{fam[:3].upper()}-{random.randint(100,999)}",
            "produit": f"{prod} {cal}", "famille": fam, "quantite": qte,
            "pu_local": pu_local, "montant_local": montant_local,
        })
    cmds.sort(key=lambda x: (x["date"], x["reference"]))
    c["commandes"] = cmds
    c["lignes"] = len(cmds)
    c["total_local"] = sum(x["montant_local"] for x in cmds)
    c["total_usd"] = round(sum(round(x["montant_local"] / TAUX[c["devise"]], 2)
                               for x in cmds), 2)
    c["total_usd_direct"] = round(c["total_local"] / TAUX[c["devise"]], 2)
    for x in cmds:
        lignes.append({
            "date": x["date"], "code": c["code"], "client": c["nom"],
            "agence": c["agence"], "pays": c["pays"], "devise": c["devise"],
            "famille": x["famille"], "produit": x["produit"],
            "quantite": x["quantite"], "pu_local": x["pu_local"],
            "montant_local": x["montant_local"], "taux": TAUX[c["devise"]],
            "montant_usd": round(x["montant_local"] / TAUX[c["devise"]], 2),
        })

lignes.sort(key=lambda x: (x["date"], x["code"], x["produit"]))
TOTAL_USD = round(sum(x["montant_usd"] for x in lignes), 2)
TOTAL_SOURCES = round(sum(c["total_usd"] for c in clients), 2)

print(f"  {len(clients)} clients · {len(lignes)} lignes · "
      f"total {TOTAL_USD:,.2f} USD")
print(f"  écart consolidé / sources : {round(TOTAL_USD - TOTAL_SOURCES, 2)}")

ECART_METHODE = round(sum(c["total_usd_direct"] for c in clients) - TOTAL_USD, 2)
print(f"  écart de méthode de conversion : {ECART_METHODE} USD")

# ══ 4 · Les 200 classeurs clients ═════════════════════════════════════
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

Ftitre  = Font(name="Calibri", size=14, bold=True)
Fentete = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
Faide   = Font(name="Calibri", size=10, italic=True, color="FF6B7280")
Fgras   = Font(name="Calibri", size=11, bold=True)
Fbleu   = Font(name="Calibri", size=11, color="FF0000FF")
Fnoir   = Font(name="Calibri", size=11, color="FF000000")
Fvert   = Font(name="Calibri", size=11, color="FF008000")
Rentete = PatternFill("solid", fgColor="FF243044")
Rjaune  = PatternFill("solid", fgColor="FFFFF6D8")
Rbleu   = PatternFill("solid", fgColor="FFEFF4FE")
Bord    = Border(bottom=Side(style="thin", color="FFD6DBE1"))
USD, USD0, LOC, PCT, INT = ('#,##0.00\\ "USD"', '#,##0\\ "USD"',
                            "#,##0", "0.00%", "#,##0")
DATE = "dd/mm/yyyy"

def entete(ws, cols, ligne=1, depart=1):
    for k, c in enumerate(cols, start=depart):
        cel = ws.cell(ligne, k, c); cel.font, cel.fill = Fentete, Rentete
        cel.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[ligne].height = 22

def larg(ws, specs, depart=1):
    for k, w in enumerate(specs, start=depart):
        ws.column_dimensions[get_column_letter(k)].width = w

def tableau(ws, nom, ref):
    t = Table(displayName=nom, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)

COLS_CLIENT = ["Date", "Reference", "Produit", "Famille", "Quantite",
               "PU_Local", "Montant_Local"]

def classeur_client(c):
    wb = openpyxl.Workbook()
    ws = wb.active
    #  Deux fichiers sur deux cents ont un nom de feuille en majuscules.
    #  C'est exactement ce qui fait planter un script écrit un vendredi.
    ws.title = "COMMANDES" if c["anomalie"].startswith("feuille") else "Commandes"
    cols = list(COLS_CLIENT)
    if c["anomalie"].startswith("colonne"):
        cols.insert(6, "Remise_Pct")
    larg(ws, [12, 12, 26, 13, 11, 14, 16, 14][:len(cols)])
    ws["A1"] = c["nom"]; ws["A1"].font = Ftitre
    ws["A2"] = (f"Code client {c['code']}  ·  agence de {c['agence']} "
                f"({c['pays']})  ·  montants en {c['devise']}")
    ws["A2"].font = Faide
    ws["A3"] = "Commandes du mois d'août 2026"; ws["A3"].font = Fgras
    entete(ws, cols, ligne=5)
    r = 6
    for x in c["commandes"]:
        vals = [x["date"], x["reference"], x["produit"], x["famille"],
                x["quantite"], x["pu_local"]]
        if c["anomalie"].startswith("colonne"):
            vals.append(0)
        vals.append(x["montant_local"])
        for j, v in enumerate(vals, start=1):
            cel = ws.cell(r, j, v)
            cel.border = Bord
            if j == 1: cel.number_format = DATE
            elif cols[j - 1] in ("Quantite",): cel.number_format = INT
            elif cols[j - 1] == "Remise_Pct": cel.number_format = PCT
            elif cols[j - 1] in ("PU_Local", "Montant_Local"):
                cel.number_format = LOC
        r += 1
    if c["commandes"]:
        tableau(ws, "t_Commandes", f"A5:{get_column_letter(len(cols))}{r-1}")
    #  Six fichiers portent une ligne TOTAL sous le tableau. Un script
    #  qui somme la colonne entière compte tout deux fois.
    if c["anomalie"].startswith("ligne TOTAL"):
        ws.cell(r + 1, 1, "TOTAL").font = Fgras
        cel = ws.cell(r + 1, len(cols), c["total_local"])
        cel.font, cel.number_format = Fgras, LOC
    nom = f"{c['code']}_{c['nom'].replace(' ', '_').replace('/', '-')}.xlsx"
    wb.save(CLIENTS_DIR / nom)
    c["fichier"] = nom

for c in clients:
    classeur_client(c)
print(f"  {len(clients)} classeurs clients écrits")

zchemin = RACINE / "J10_Clients.zip"
with zipfile.ZipFile(zchemin, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
    for f in sorted(CLIENTS_DIR.iterdir()):
        z.write(f, f"J10_Clients/{f.name}")
print(f"  J10_Clients.zip  {zchemin.stat().st_size/1e6:.1f} Mo")

# ══ 5 · Les douze rapports du lundi matin ═════════════════════════════
#  Six agences et six familles : deux découpes du même total, qui
#  doivent toutes deux se re-sommer au consolidé. C'est ce double
#  contrôle qui prouve que les douze PDF sont justes.
RAPPORTS = ([("Agence", a[0]) for a in AGENCES]
            + [("Famille", f) for f in FAMILLES])
MOIS_ETIQ = "2026-08"

def sans_accent(t):
    tr = str.maketrans("àâäéèêëîïôöùûüçÀÂÄÉÈÊËÎÏÔÖÙÛÜÇ'’ ",
                       "aaaeeeeiioouuucAAAEEEEIIOOUUUC___")
    return t.translate(tr)

for t, v in RAPPORTS:
    pass
RAPPORTS = [(t, v, f"{MOIS_ETIQ}_{t}_{sans_accent(v)}.pdf") for t, v in RAPPORTS]

def total_rapport(t, v):
    cle = "agence" if t == "Agence" else "famille"
    return round(sum(x["montant_usd"] for x in lignes if x[cle] == v), 2)

TOTAUX_RAPPORTS = {v: total_rapport(t, v) for t, v, _f in RAPPORTS}
LIGNES_RAPPORTS = {v: sum(1 for x in lignes
                          if x["agence" if t == "Agence" else "famille"] == v)
                   for t, v, _f in RAPPORTS}

R = {
 "clients": len(clients),
 "clients_actifs": sum(1 for c in clients if c["lignes"] > 0),
 "classeurs_vides": sum(1 for c in clients if c["lignes"] == 0),
 "classeurs_anomalie": sum(1 for c in clients if c["anomalie"]),
 "lignes": len(lignes),
 "total_usd": TOTAL_USD,
 "ecart_methode": ECART_METHODE,
 "rapports": len(RAPPORTS),
 "total_agences": round(sum(TOTAUX_RAPPORTS[a[0]] for a in AGENCES), 2),
 "total_familles": round(sum(TOTAUX_RAPPORTS[f] for f in FAMILLES), 2),
 "plus_gros_client_usd": round(max(c["total_usd"] for c in clients), 2),
 "plus_gros_client": max(clients, key=lambda c: c["total_usd"])["nom"],
 "clients_20k": sum(1 for c in clients if c["total_usd"] > 20_000),
 "totaux_rapports": TOTAUX_RAPPORTS,
 "taille_zip_mo": round(zchemin.stat().st_size / 1e6, 2),
}
for k, v in R.items():
    if not isinstance(v, dict): print(f"  {k:<22} {v}")

# ══ 6 · Le classeur du TP ═════════════════════════════════════════════
def feuille_sources(wb):
    ws = wb.create_sheet("SOURCES")
    larg(ws, [34, 13, 30, 14, 9, 9, 16, 15, 17, 30])
    ws["A1"] = "Les 200 classeurs reçus — un récapitulatif par fichier"
    ws["A1"].font = Ftitre
    ws["A2"] = ("Produit par la consolidation. C'est cette feuille qui permet "
                "de réconcilier : si elle ne tombe pas d'accord avec DONNEES, "
                "un fichier a été lu deux fois ou pas du tout.")
    ws["A2"].font = Faide
    cols = ["Fichier", "Code_Client", "Client", "Agence", "Devise", "Lignes",
            "Total_Local", "Total_USD", "Total_USD_direct", "Anomalie"]
    entete(ws, cols, ligne=4)
    for k, c in enumerate(clients, start=5):
        for j, v in enumerate([c["fichier"], c["code"], c["nom"], c["agence"],
                               c["devise"], c["lignes"], c["total_local"],
                               c["total_usd"], c["total_usd_direct"],
                               c["anomalie"] or None], start=1):
            cel = ws.cell(k, j, v); cel.border = Bord
            if j == 6: cel.number_format = INT
            elif j == 7: cel.number_format = LOC
            elif j in (8, 9): cel.number_format = USD
    tableau(ws, "t_Sources", f"A4:J{4+len(clients)}")
    return ws

def feuille_donnees(wb):
    ws = wb.create_sheet("DONNEES")
    larg(ws, [11, 12, 30, 13, 12, 26, 10, 13, 15, 8, 11, 14])
    ws["A1"] = "La consolidation — une ligne par ligne de commande"
    ws["A1"].font = Ftitre
    ws["A2"] = (f"{len(lignes)} lignes, août 2026. Conversion ligne à ligne, "
                "au taux du mois de la devise de l'agence.")
    ws["A2"].font = Faide
    cols = ["Date", "Code_Client", "Client", "Agence", "Famille", "Produit",
            "Quantite", "PU_Local", "Montant_Local", "Devise", "Taux",
            "Montant_USD"]
    entete(ws, cols, ligne=4)
    for k, x in enumerate(lignes, start=5):
        for j, v in enumerate([x["date"], x["code"], x["client"], x["agence"],
                               x["famille"], x["produit"], x["quantite"],
                               x["pu_local"], x["montant_local"], x["devise"],
                               x["taux"], x["montant_usd"]], start=1):
            cel = ws.cell(k, j, v)
            if j == 1: cel.number_format = DATE
            elif j == 7: cel.number_format = INT
            elif j in (8, 9): cel.number_format = LOC
            elif j == 11: cel.number_format = "#,##0.00"
            elif j == 12: cel.number_format = USD
    tableau(ws, "t_Ventes", f"A4:L{4+len(lignes)}")
    ws.freeze_panes = "A5"
    return ws

def feuille_parametres(wb, corrige):
    ws = wb.create_sheet("PARAMETRES")
    larg(ws, [16, 22, 42, 18, 12, 40])
    ws["A1"] = "Ce que la macro doit produire"; ws["A1"].font = Ftitre
    ws["A2"] = ("Douze rapports : six agences et six familles. Deux découpes "
                "du même total — et chacune doit se re-sommer au consolidé.")
    ws["A2"].font = Faide
    entete(ws, ["Type", "Valeur", "Nom_Fichier", "Total_USD"], ligne=4)
    for k, (t, v, f) in enumerate(RAPPORTS, start=5):
        ws.cell(k, 1, t).border = Bord
        ws.cell(k, 2, v).border = Bord
        ws.cell(k, 3, f).border = Bord
        c1 = ws.cell(k, 4); c1.fill, c1.border, c1.number_format = Rjaune, Bord, USD
        if corrige:
            #  SOMME.SI.ENS accepte « * » comme joker : le critère devient
            #  « n'importe quel texte », ce qui neutralise la découpe
            #  qui ne s'applique pas à cette ligne.
            c1.value = (f'=SUMIFS(t_Ventes[Montant_USD],'
                        f't_Ventes[Agence],IF($A{k}="Agence",$B{k},"*"),'
                        f't_Ventes[Famille],IF($A{k}="Famille",$B{k},"*"))')
            c1.font = Fnoir
    tableau(ws, "t_Rapports", f"A4:D{4+len(RAPPORTS)}")

    ws["A19"] = "Les paramètres d'exécution"; ws["A19"].font = Ftitre
    for k, (lib, val, aide) in enumerate([
        ("Dossier de sortie", "Rapports",
         "un sous-dossier du dossier de ce classeur — jamais un chemin absolu"),
        ("Période", MOIS_ETIQ, "sert à nommer les fichiers"),
        ("Taux CDF", TAUX["CDF"], "figé pour le mois"),
        ("Taux XOF", TAUX["XOF"], "arrimé à l'euro"),
        ("Taux XAF", TAUX["XAF"], "même parité que le XOF — ce n'est pas une coquille"),
    ], start=21):
        ws.cell(k, 1, lib).font = Fgras
        c = ws.cell(k, 2, val); c.font, c.fill, c.border = Fbleu, Rbleu, Bord
        if isinstance(val, float): c.number_format = "#,##0.00"
        ws.cell(k, 3, aide).font = Faide
    return ws

# ── RAPPORT · la page que la macro exporte, douze fois ────────────────
def feuille_rapport(wb, corrige):
    ws = wb.create_sheet("RAPPORT")
    larg(ws, [30, 20, 16, 14, 40])
    ws["A1"] = "BAOBAB — rapport de synthèse"; ws["A1"].font = Ftitre
    ws["A2"] = "Août 2026"; ws["A2"].font = Faide
    ws["A4"] = "Découpe"; ws["A4"].font = Fgras
    ws["A5"] = "Valeur"; ws["A5"].font = Fgras
    for ref, val in (("B4", "Agence"), ("B5", AGENCES[0][0])):
        c = ws[ref]; c.font, c.fill, c.border = Fbleu, Rbleu, Bord
        if corrige: c.value = val
    ws["C4"] = "← ces deux cellules sont écrites par la macro, à chaque tour"
    ws["C4"].font = Faide
    KPI = [("Chiffre d'affaires", USD,
            '=SUMIFS(t_Ventes[Montant_USD],t_Ventes[Agence],'
            'IF($B$4="Agence",$B$5,"*"),t_Ventes[Famille],'
            'IF($B$4="Famille",$B$5,"*"))'),
           ("Lignes de commande", INT,
            '=COUNTIFS(t_Ventes[Agence],IF($B$4="Agence",$B$5,"*"),'
            't_Ventes[Famille],IF($B$4="Famille",$B$5,"*"))'),
           ("Quantité totale", INT,
            '=SUMIFS(t_Ventes[Quantite],t_Ventes[Agence],'
            'IF($B$4="Agence",$B$5,"*"),t_Ventes[Famille],'
            'IF($B$4="Famille",$B$5,"*"))'),
           ("Panier moyen par ligne", USD, "=IF(B8=0,0,B7/B8)"),
           ("Part du consolidé", PCT, "=IF(REPONSES!$B$15=0,0,B7/REPONSES!$B$15)")]
    for k, (lib, fmt, f) in enumerate(KPI, start=7):
        ws.cell(k, 1, lib).font = Fgras
        c = ws.cell(k, 2); c.number_format, c.border = fmt, Bord
        if corrige: c.value, c.font = f, Fnoir
    ws["A13"] = "Le détail"; ws["A13"].font = Ftitre
    entete(ws, ["Ligne", "Chiffre d'affaires", "Lignes", "Part"], ligne=15)
    #  Le détail bascule : une découpe par agence se détaille par famille,
    #  et inversement. Une seule page sert donc aux douze rapports.
    for k in range(16, 22):
        i = k - 16
        c0 = ws.cell(k, 1); c0.border = Bord
        c1 = ws.cell(k, 2); c1.number_format, c1.border = USD, Bord
        c2 = ws.cell(k, 3); c2.number_format, c2.border = INT, Bord
        c3 = ws.cell(k, 4); c3.number_format, c3.border = PCT, Bord
        if corrige:
            c0.value = (f'=IF($B$4="Agence",INDEX(PARAMETRES!$B$11:$B$16,{i+1}),'
                        f'INDEX(PARAMETRES!$B$5:$B$10,{i+1}))')
            c1.value = (f'=SUMIFS(t_Ventes[Montant_USD],t_Ventes[Agence],'
                        f'IF($B$4="Agence",$B$5,$A{k}),t_Ventes[Famille],'
                        f'IF($B$4="Agence",$A{k},$B$5))')
            c2.value = (f'=COUNTIFS(t_Ventes[Agence],'
                        f'IF($B$4="Agence",$B$5,$A{k}),t_Ventes[Famille],'
                        f'IF($B$4="Agence",$A{k},$B$5))')
            c3.value = f"=IF($B$7=0,0,B{k}/$B$7)"
            for c in (c0, c1, c2, c3): c.font = Fnoir
    ws.cell(22, 1, "Total").font = Fgras
    if corrige:
        c = ws.cell(22, 2, "=SUM(B16:B21)"); c.font, c.number_format = Fgras, USD
        c = ws.cell(22, 3, "=SUM(C16:C21)"); c.font, c.number_format = Fgras, INT
    ws.print_area = "A1:D22"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws

# ── JOURNAL · ce que la macro écrit en tournant ───────────────────────
def feuille_journal(wb, corrige):
    ws = wb.create_sheet("JOURNAL")
    larg(ws, [20, 44, 12, 12, 18, 40])
    ws["A1"] = "Le journal d'exécution"; ws["A1"].font = Ftitre
    ws["A2"] = ("Une ligne par rapport produit. Une macro qui ne laisse pas de "
                "trace est une macro dont on ne peut rien dire le lendemain.")
    ws["A2"].font = Faide
    entete(ws, ["Horodatage", "Fichier", "Statut", "Lignes", "Total_USD",
                "Message"], ligne=4)
    n = len(RAPPORTS) if corrige else 1
    base = dt.datetime(2026, 9, 21, 7, 42, 11)
    for k in range(5, 5 + n):
        for j in range(1, 7):
            ws.cell(k, j).border = Bord
        if corrige:
            i = k - 5
            t, v, f = RAPPORTS[i]
            ws.cell(k, 1, base + dt.timedelta(seconds=i * 3 + 2)).number_format = \
                "dd/mm/yyyy hh:mm:ss"
            ws.cell(k, 2, f)
            ws.cell(k, 3, "OK")
            ws.cell(k, 4, LIGNES_RAPPORTS[v]).number_format = INT
            ws.cell(k, 5, TOTAUX_RAPPORTS[v]).number_format = USD
            ws.cell(k, 6, "")
    tableau(ws, "t_Journal", f"A4:F{4+n}")
    return ws

# ── REPONSES ──────────────────────────────────────────────────────────
LUS = [
 ("Durée d'exécution de la macro, en secondes", "0.0", None),
 ("Nombre de PDF réellement présents dans le dossier de sortie", INT, 12),
]
REPONSES = [
 ("Nombre de classeurs clients reçus", "=ROWS(t_Sources[Fichier])", INT,
  R["clients"]),
 ("Nombre de classeurs vides", "=COUNTIF(t_Sources[Lignes],0)", INT,
  R["classeurs_vides"]),
 ("Nombre de classeurs porteurs d'une anomalie",
  '=COUNTIF(t_Sources[Anomalie],"?*")', INT, R["classeurs_anomalie"]),
 ("Nombre de clients ayant commandé", '=COUNTIF(t_Sources[Lignes],">0")', INT,
  R["clients_actifs"]),
 ("Nombre de lignes consolidées", "=ROWS(t_Ventes[Date])", INT, R["lignes"]),
 ("Total consolidé, en USD", "=SUM(t_Ventes[Montant_USD])", USD,
  R["total_usd"]),
 ("R_ECART_SOURCES — consolidé moins récapitulatif  ·  doit valoir 0",
  "=ROUND(B15-SUM(t_Sources[Total_USD]),2)", USD, 0),
 ("Écart entre les deux méthodes de conversion",
  "=ROUND(SUM(t_Sources[Total_USD_direct])-SUM(t_Sources[Total_USD]),2)",
  USD, R["ecart_methode"]),
 ("Nombre de rapports à produire", "=ROWS(t_Rapports[Nom_Fichier])", INT,
  R["rapports"]),
 ("Somme des six rapports d'agence",
  '=SUMIF(t_Rapports[Type],"Agence",t_Rapports[Total_USD])', USD,
  R["total_agences"]),
 ("R_ECART_AGENCES — la découpe par agence  ·  doit valoir 0",
  "=ROUND(B19-B15,2)", USD, 0),
 ("R_ECART_FAMILLES — la découpe par famille  ·  doit valoir 0",
  '=ROUND(SUMIF(t_Rapports[Type],"Famille",t_Rapports[Total_USD])-B15,2)',
  USD, 0),
 ("Rapports effectivement produits, d'après le journal",
  '=COUNTIF(t_Journal[Statut],"OK")', INT, R["rapports"]),
 ("Le plus gros client du mois, en USD", "=MAX(t_Sources[Total_USD])", USD,
  R["plus_gros_client_usd"]),
 ("Son nom",
  "=INDEX(t_Sources[Client],MATCH(MAX(t_Sources[Total_USD]),"
  "t_Sources[Total_USD],0))", None, R["plus_gros_client"]),
 ("Clients au-dessus de 20 000 USD sur le mois",
  '=COUNTIF(t_Sources[Total_USD],">20000")', INT, R["clients_20k"]),
 ("R_CONTROLE — quantité × prix moins le montant  ·  doit valoir 0",
  "=ROUND(SUMPRODUCT(t_Ventes[Quantite],t_Ventes[PU_Local])"
  "-SUM(t_Ventes[Montant_Local]),0)", LOC, 0),
]

def feuille_reponses(wb, corrige):
    ws = wb.create_sheet("REPONSES")
    larg(ws, [72, 26])
    ws["A1"] = "TP 10 — Le lundi matin d'Aïcha"; ws["A1"].font = Ftitre
    ws["A2"] = ("Les deux premières se lisent après avoir lancé la macro. "
                "Les autres se calculent, et elles se calculent même si la "
                "macro n'a pas encore tourné.")
    ws["A2"].font = Faide
    ws["A4"] = "①  CE QUE VOUS LISEZ APRÈS EXÉCUTION"; ws["A4"].font = Fgras
    for k, (lib, fmt, val) in enumerate(LUS, start=5):
        ws.cell(k, 1, lib).border = Bord
        c = ws.cell(k, 2); c.fill, c.border = Rjaune, Bord
        if fmt: c.number_format = fmt
        if corrige: c.value, c.font = (val if val is not None else 37.4), Fbleu
    ws["A8"] = "②  CE QUE LES FORMULES VÉRIFIENT"; ws["A8"].font = Fgras
    for k, (lib, f, fmt, _v) in enumerate(REPONSES, start=10):
        ws.cell(k, 1, lib).border = Bord
        c = ws.cell(k, 2); c.fill, c.border = Rjaune, Bord
        if fmt: c.number_format = fmt
        if corrige: c.value, c.font = f, Fnoir
    return ws

QUALITE = [
 ("Classeurs manquants  ·  doit valoir 0", "=200-ROWS(t_Sources[Fichier])",
  INT, 0, "deux cents fichiers reçus, deux cents fichiers lus"),
 ("Fichiers vides portant quand même un total  ·  doit valoir 0",
  '=COUNTIFS(t_Sources[Lignes],0,t_Sources[Total_USD],"<>0")', INT, 0,
  "un fichier sans ligne doit totaliser zéro, pas « rien »"),
 ("Montants négatifs dans le consolidé  ·  doit valoir 0",
  "=SUMPRODUCT((t_Ventes[Montant_USD]<0)*1)", INT, 0,
  "un avoir existe, mais il ne se glisse pas ici sans qu'on le sache"),
 ("Lignes sans agence  ·  doit valoir 0",
  '=ROWS(t_Ventes[Agence])-COUNTIF(t_Ventes[Agence],"?*")', INT, 0,
  "une ligne orpheline ne sortira dans aucun des douze rapports"),
 ("Rapports attendus manquants  ·  doit valoir 0",
  "=12-ROWS(t_Rapports[Nom_Fichier])", INT, 0,
  "six agences et six familles"),
 ("Le journal contient-il un échec ?",
  '=IF(COUNTIF(t_Journal[Statut],"KO")=0,"PASS","FAIL")', None, "PASS",
  "un seul KO interdit l'envoi — même si onze PDF sont justes"),
 ("La réconciliation passe-t-elle ?", '=IF(REPONSES!$B$16=0,"PASS","FAIL")',
  None, "PASS", "c'est le contrôle qui vaut tous les autres"),
]

def feuille_qualite(wb, corrige):
    ws = wb.create_sheet("QUALITE")
    larg(ws, [56, 18, 66])
    ws["A1"] = "Contrôles de qualité"; ws["A1"].font = Ftitre
    ws["A2"] = ("Sept contrôles, à passer avant d'envoyer quoi que ce soit. "
                "Ils se relancent tout seuls le mois prochain.")
    ws["A2"].font = Faide
    entete(ws, ["Contrôle", "Valeur", "Ce que ça veut dire"], ligne=4)
    for k, (lib, f, fmt, _v, quoi) in enumerate(QUALITE, start=5):
        ws.cell(k, 1, lib).border = Bord
        c = ws.cell(k, 2); c.fill, c.border = Rjaune, Bord
        if fmt: c.number_format = fmt
        if corrige: c.value, c.font = f, Fnoir
        ws.cell(k, 3, quoi).font = Faide
    return ws

# ── README · ANNEXE_IA ────────────────────────────────────────────────
TEXTES = {
 "README": [
  ("À quoi sert ce classeur",
   "Produire les douze rapports mensuels de BAOBAB — six agences et six "
   "familles de produits — à partir des commandes des 200 clients grands "
   "comptes. Un clic, douze PDF nommés, un journal d'exécution."),
  ("D'où viennent les données",
   "200 classeurs clients au format identique, reçus par messagerie et "
   "déposés dans le dossier J10_Clients. Consolidés dans DONNEES, "
   "récapitulés fichier par fichier dans SOURCES. Période : août 2026."),
  ("Les quinze fichiers qui ne sont pas au format",
   "Quatre sont vides — le client n'a rien commandé ce mois-ci. Six "
   "portent une ligne TOTAL sous le tableau, qu'il ne faut surtout pas "
   "additionner. Trois ont une colonne Remise_Pct en plus. Deux ont une "
   "feuille nommée COMMANDES au lieu de Commandes. La colonne Anomalie "
   "de SOURCES les liste tous les quinze."),
  ("Les taux de change",
   "Figés pour le mois : 3 145,80 CDF et 648,20 XOF comme XAF pour "
   "1 USD. Le franc CFA Ouest et le franc CFA Central sont arrimés à "
   "l'euro à la même parité : leur taux face au dollar est identique. "
   "Ce n'est pas une coquille."),
  ("Comment on convertit, et pourquoi ça compte",
   "Ligne à ligne, jamais sur le total. Les deux méthodes sont "
   "défendables ; elles écartent de 0,54 USD sur ce mois. Ce qui n'est "
   "pas défendable, c'est de ne pas savoir laquelle on a utilisée."),
  ("Comment lancer la production",
   "Bouton « Produire les douze rapports » sur la feuille PARAMETRES. "
   "Le dossier de sortie est un sous-dossier du dossier de ce classeur, "
   "nommé dans PARAMETRES!B21. Durée observée : environ 40 secondes."),
  ("Ce qu'il faut faire si ça casse",
   "Lire la colonne Message du JOURNAL : elle nomme le rapport qui a "
   "échoué et la raison. Les trois causes vues jusqu'ici : le dossier de "
   "sortie n'existe pas, un PDF du même nom est ouvert dans un lecteur, "
   "et le classeur a été ouvert depuis la messagerie sans avoir été "
   "enregistré ailleurs d'abord."),
  ("Ce que ce classeur ne fait pas",
   "Il n'envoie aucun courriel, il ne va pas chercher les fichiers "
   "clients tout seul, et il ne détecte pas un client qui aurait oublié "
   "d'envoyer son fichier. Trois manques connus, écrits ici."),
  ("Livraison",
   "En .zip, jamais en .xlsm nu : les pièces jointes .xlsm sont bloquées "
   "par la plupart des messageries, et le Mark of the Web désactive les "
   "macros d'un fichier venu d'Internet. La procédure de déblocage est "
   "dans le fichier LISEZ-MOI.txt du zip."),
  ("Journal des versions", "v1.0 — septembre 2026."),
 ],
 "ANNEXE_IA": [
  ("Le prompt qui a produit la macro",
   "Ce que vous avez demandé, mot pour mot."),
  ("Ce que l'IA a écrit et qui était faux",
   "Décrivez précisément l'erreur, et comment vous l'avez vue."),
  ("Le V qui l'a attrapée",
   "V1 source · V2 calcul · V3 cohérence · V4 terrain · V5 lisibilité."),
  ("Ce que vous avez corrigé dans le code",
   "Ligne par ligne, ce qui a changé et pourquoi."),
  ("Ce que vous n'avez pas suivi",
   "Une suggestion de l'IA que vous avez écartée, et pourquoi."),
 ],
}

def feuille_texte(wb, nom, corrige):
    ws = wb.create_sheet(nom)
    larg(ws, [32, 102])
    ws["A1"] = nom; ws["A1"].font = Ftitre
    for k, (cle, val) in enumerate(TEXTES[nom], start=3):
        c = ws.cell(k, 1, cle); c.font, c.border = Fgras, Bord
        d = ws.cell(k, 2, val if corrige else ""); d.border = Bord
        d.alignment = Alignment(wrap_text=True, vertical="top")
        if not corrige: d.fill = Rjaune
        ws.row_dimensions[k].height = 48
    return ws

PLAGES = [
 ("R_ECART_SOURCES",  "REPONSES!$B$16"),
 ("R_ECART_AGENCES",  "REPONSES!$B$20"),
 ("R_ECART_FAMILLES", "REPONSES!$B$21"),
 ("R_CONTROLE",       "REPONSES!$B$26"),
 ("P_TYPE",           "RAPPORT!$B$4"),
 ("P_VALEUR",         "RAPPORT!$B$5"),
 ("P_DOSSIER",        "PARAMETRES!$B$21"),
]
#  RAPPORT et JOURNAL sont des constructions libres : la mise en page du
#  rapport appartient à l'apprenant, et le journal est écrit par sa macro.
ZONES_LIBRES = [
 ("LIBRE_RAPPORT", "RAPPORT!$A$1:$E$25"),
 ("LIBRE_JOURNAL", "JOURNAL!$A$1:$F$20"),
]

def classeur_tp(corrige):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    feuille_texte(wb, "README", corrige)
    feuille_sources(wb)
    feuille_donnees(wb)
    feuille_parametres(wb, corrige)
    feuille_rapport(wb, corrige)
    feuille_journal(wb, corrige)
    feuille_qualite(wb, corrige)
    feuille_reponses(wb, corrige)
    feuille_texte(wb, "ANNEXE_IA", corrige)
    if corrige:
        for nom, cible in PLAGES + ZONES_LIBRES:
            wb.defined_names[nom] = DefinedName(nom, attr_text=cible)
    nom = "TP10_CORRIGE.xlsx" if corrige else "TP10_DEPART.xlsx"
    wb.save(RACINE / nom)
    print(f"  {nom}  {(RACINE/nom).stat().st_size/1e6:.2f} Mo")

classeur_tp(False)
classeur_tp(True)
json.dump(R, open(RACINE / "REFERENCES_M10.json", "w"), indent=1, ensure_ascii=False)
print("  REFERENCES_M10.json")
